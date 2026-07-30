from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import func, select

from app.models.credit_payment import CreditPayment
from app.models.customer import Customer
from app.models.operating_expense import OperatingExpense
from app.models.operating_expense_category import OperatingExpenseCategory
from app.models.product import Product
from app.models.sale import Sale
from app.models.sale_item import SaleItem
from app.models.user import User
from app.services.excel_export import EXCEL_MEDIA_TYPE

CUSTOMER_PERIOD_PARAMS = {
    "start_date": "2026-07-01",
    "end_date": "2026-07-31",
}
ANALYTICS_PARAMS = {
    "start_date": "2026-07-10",
    "end_date": "2026-07-12",
}


def test_customer_sales_summary_without_sales_returns_zeroes(
    api_client,
    api_auth_headers,
) -> None:
    client, testing_session_local = api_client
    customer_id = _create_empty_customer(testing_session_local)

    response = client.get(
        f"/api/v1/customers/{customer_id}/sales-summary",
        headers=api_auth_headers,
    )

    assert response.status_code == 200
    body = response.json()
    assert body["cash_sales_count"] == 0
    assert body["cash_sales_total"] == "0.00"
    assert body["credit_sales_count"] == 0
    assert body["credit_sales_total"] == "0.00"
    assert body["total_sales_count"] == 0
    assert body["total_sales_amount"] == "0.00"
    assert body["total_collected_amount"] == "0.00"
    assert body["pending_credit_amount"] == "0.00"
    assert body["last_sale_date"] is None


def test_customer_sales_summary_calculates_cash_credit_and_balance(
    api_client,
    api_auth_headers,
) -> None:
    client, testing_session_local = api_client
    seeded = _seed_reporting_data(testing_session_local)

    response = client.get(
        f"/api/v1/customers/{seeded['customer_a_id']}/sales-summary",
        headers=api_auth_headers,
        params=CUSTOMER_PERIOD_PARAMS,
    )

    assert response.status_code == 200
    body = response.json()
    assert body["cash_sales_count"] == 1
    assert body["cash_sales_total"] == "100.00"
    assert body["credit_sales_count"] == 1
    assert body["credit_sales_total"] == "60.00"
    assert body["total_sales_count"] == 2
    assert body["total_sales_amount"] == "160.00"
    assert body["total_collected_amount"] == "20.00"
    assert body["pending_credit_amount"] == "40.00"
    assert body["last_sale_date"] == "2026-07-11"


def test_customer_sales_summary_excludes_voided_sales_from_main_totals(
    api_client,
    api_auth_headers,
) -> None:
    client, testing_session_local = api_client
    seeded = _seed_reporting_data(testing_session_local)

    response = client.get(
        f"/api/v1/customers/{seeded['customer_a_id']}/sales-summary",
        headers=api_auth_headers,
        params=CUSTOMER_PERIOD_PARAMS,
    )

    body = response.json()
    assert body["total_sales_amount"] == "160.00"
    assert body["cancelled_sales_count"] == 1
    assert body["cancelled_sales_total"] == "999.00"


def test_customer_sales_summary_respects_date_filters(
    api_client,
    api_auth_headers,
) -> None:
    client, testing_session_local = api_client
    seeded = _seed_reporting_data(testing_session_local)

    response = client.get(
        f"/api/v1/customers/{seeded['customer_a_id']}/sales-summary",
        headers=api_auth_headers,
        params={"start_date": "2026-07-11", "end_date": "2026-07-11"},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["cash_sales_count"] == 0
    assert body["credit_sales_count"] == 1
    assert body["total_sales_amount"] == "60.00"
    assert body["period_start"] == "2026-07-11"
    assert body["period_end"] == "2026-07-11"


def test_customer_sales_detail_returns_items(
    api_client,
    api_auth_headers,
) -> None:
    client, testing_session_local = api_client
    seeded = _seed_reporting_data(testing_session_local)

    response = client.get(
        f"/api/v1/customers/{seeded['customer_a_id']}/sales",
        headers=api_auth_headers,
        params={
            **CUSTOMER_PERIOD_PARAMS,
            "sale_type": "credit",
            "status": "active",
        },
    )

    assert response.status_code == 200
    sales = response.json()
    assert len(sales) == 1
    assert sales[0]["sale_number"] == "VTA-A-CREDIT"
    assert sales[0]["items"][0]["product_name_snapshot"] == "Bidon 20L"
    assert sales[0]["items"][0]["subtotal"] == "60.00"


def test_operator_can_read_customer_sales_summary(
    api_client,
    api_operator_auth_headers,
) -> None:
    client, testing_session_local = api_client
    customer_id = _create_empty_customer(testing_session_local)

    response = client.get(
        f"/api/v1/customers/{customer_id}/sales-summary",
        headers=api_operator_auth_headers,
    )

    assert response.status_code == 200


def test_customer_credit_export_requires_customer(
    api_client,
    api_auth_headers,
) -> None:
    client, _ = api_client

    response = client.get(
        "/api/v1/reports/exports/customer-credit-sales.xlsx",
        headers=api_auth_headers,
    )

    assert response.status_code == 422


def test_customer_credit_export_blocks_operator(
    api_client,
    api_operator_auth_headers,
) -> None:
    client, _ = api_client

    response = client.get(
        (
            "/api/v1/reports/exports/customer-credit-sales.xlsx"
            "?customer_id=1"
        ),
        headers=api_operator_auth_headers,
    )

    assert response.status_code == 403


def test_customer_credit_export_returns_detailed_valid_xlsx_for_admin(
    api_client,
    api_auth_headers,
) -> None:
    client, testing_session_local = api_client
    seeded = _seed_reporting_data(testing_session_local)

    response = client.get(
        "/api/v1/reports/exports/customer-credit-sales.xlsx",
        headers=api_auth_headers,
        params={
            "customer_id": seeded["customer_a_id"],
            **CUSTOMER_PERIOD_PARAMS,
        },
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == EXCEL_MEDIA_TYPE
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert workbook.sheetnames == ["Resumen", "Detalle ventas crédito"]
    summary = workbook["Resumen"]
    assert summary["B2"].value == "Cliente A"
    assert summary["B5"].value == 60
    assert summary["B6"].value == 20
    assert summary["B7"].value == 40
    assert summary["B8"].value == 1
    detail = workbook["Detalle ventas crédito"]
    assert detail["B2"].value == seeded["credit_sale_id"]
    assert detail["E2"].value == "Bidon 20L"
    assert detail["H2"].value == 60
    workbook.close()


def test_customer_credit_export_without_sales_contains_clear_message(
    api_client,
    api_auth_headers,
) -> None:
    client, testing_session_local = api_client
    customer_id = _create_empty_customer(testing_session_local)

    response = client.get(
        "/api/v1/reports/exports/customer-credit-sales.xlsx",
        headers=api_auth_headers,
        params={"customer_id": customer_id},
    )

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert (
        workbook["Detalle ventas crédito"]["A2"].value
        == "Sin ventas a crédito en el periodo"
    )
    workbook.close()


def test_admin_analytics_calculates_sales_mix_collections_and_expenses(
    api_client,
    api_auth_headers,
) -> None:
    client, testing_session_local = api_client
    _seed_reporting_data(testing_session_local)

    response = client.get(
        "/api/v1/reports/admin-dashboard-analytics",
        headers=api_auth_headers,
        params=ANALYTICS_PARAMS,
    )

    assert response.status_code == 200
    body = response.json()
    assert body["total_sales_amount"] == "410.00"
    assert body["cash_sales_amount"] == "150.00"
    assert body["credit_sales_amount"] == "260.00"
    assert body["collections_amount"] == "20.00"
    assert body["pending_credit_amount"] == "240.00"
    assert body["expenses_amount"] == "35.00"
    assert body["net_amount"] is None
    assert body["net_amount_note"]
    assert body["customers_served_count"] == 2
    assert body["sales_mix"] == [
        {"sale_type": "cash", "amount": "150.00", "percentage": "36.59"},
        {"sale_type": "credit", "amount": "260.00", "percentage": "63.41"},
    ]


def test_admin_analytics_calculates_best_day_average_daily_and_top_customers(
    api_client,
    api_auth_headers,
) -> None:
    client, testing_session_local = api_client
    _seed_reporting_data(testing_session_local)

    response = client.get(
        "/api/v1/reports/admin-dashboard-analytics",
        headers=api_auth_headers,
        params=ANALYTICS_PARAMS,
    )

    body = response.json()
    assert body["best_sales_day"] == {
        "date": "2026-07-11",
        "amount": "260.00",
    }
    assert body["average_daily_sales"] == "136.67"
    assert body["top_customers"][0] == {
        "customer_id": body["top_customers"][0]["customer_id"],
        "customer_name": "Cliente B",
        "sales_count": 1,
        "total_amount": "200.00",
    }
    assert [row["total_sales_amount"] for row in body["daily_sales"]] == [
        "100.00",
        "260.00",
        "50.00",
    ]


def test_admin_analytics_respects_dates_and_excludes_voided_records(
    api_client,
    api_auth_headers,
) -> None:
    client, testing_session_local = api_client
    _seed_reporting_data(testing_session_local)

    response = client.get(
        "/api/v1/reports/admin-dashboard-analytics",
        headers=api_auth_headers,
        params={"start_date": "2026-07-10", "end_date": "2026-07-10"},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["total_sales_amount"] == "100.00"
    assert body["credit_sales_amount"] == "0.00"
    assert body["average_daily_sales"] == "100.00"
    assert body["best_sales_day"]["date"] == "2026-07-10"


def test_admin_analytics_blocks_operator(
    api_client,
    api_operator_auth_headers,
) -> None:
    client, _ = api_client

    response = client.get(
        "/api/v1/reports/admin-dashboard-analytics",
        headers=api_operator_auth_headers,
    )

    assert response.status_code == 403


def test_block_62a_read_endpoints_do_not_modify_business_data(
    api_client,
    api_auth_headers,
) -> None:
    client, testing_session_local = api_client
    seeded = _seed_reporting_data(testing_session_local)
    before = _business_counts(testing_session_local)

    summary_response = client.get(
        f"/api/v1/customers/{seeded['customer_a_id']}/sales-summary",
        headers=api_auth_headers,
        params=CUSTOMER_PERIOD_PARAMS,
    )
    detail_response = client.get(
        f"/api/v1/customers/{seeded['customer_a_id']}/sales",
        headers=api_auth_headers,
        params=CUSTOMER_PERIOD_PARAMS,
    )
    analytics_response = client.get(
        "/api/v1/reports/admin-dashboard-analytics",
        headers=api_auth_headers,
        params=ANALYTICS_PARAMS,
    )
    export_response = client.get(
        "/api/v1/reports/exports/customer-credit-sales.xlsx",
        headers=api_auth_headers,
        params={"customer_id": seeded["customer_a_id"]},
    )

    assert {
        summary_response.status_code,
        detail_response.status_code,
        analytics_response.status_code,
        export_response.status_code,
    } == {200}
    assert _business_counts(testing_session_local) == before


def _create_empty_customer(testing_session_local) -> int:
    with testing_session_local() as session:
        customer = Customer(name="Cliente sin ventas")
        session.add(customer)
        session.commit()
        return customer.id


def _seed_reporting_data(testing_session_local) -> dict[str, int]:
    voided_at = datetime(2026, 7, 12, 18, 0, tzinfo=timezone.utc)
    with testing_session_local() as session:
        user = session.scalar(select(User).where(User.username.in_(("admin", "operator"))))
        assert user is not None
        customer_a = Customer(
            name="Cliente A",
            credit_allowed=True,
            credit_limit=Decimal("500.00"),
        )
        customer_b = Customer(
            name="Cliente B",
            credit_allowed=True,
            credit_limit=Decimal("500.00"),
        )
        product = Product(
            sku="B20",
            name="Bidon 20L",
            category="water",
            unit="unit",
            default_price=Decimal("50.00"),
            cost_estimate=Decimal("20.00"),
            product_kind="finished_good",
        )
        expense_category = OperatingExpenseCategory(name="Combustible")
        session.add_all([customer_a, customer_b, product, expense_category])
        session.flush()

        sales = [
            Sale(
                sale_number="VTA-A-CASH",
                sale_date=date(2026, 7, 10),
                customer_id=customer_a.id,
                customer_name_snapshot=customer_a.name,
                sale_type="cash",
                payment_method="cash",
                status="active",
                total_amount=Decimal("100.00"),
                paid_amount=Decimal("100.00"),
                balance_amount=Decimal("0.00"),
                credit_status="none",
                created_by=user.id,
            ),
            Sale(
                sale_number="VTA-A-CREDIT",
                sale_date=date(2026, 7, 11),
                customer_id=customer_a.id,
                customer_name_snapshot=customer_a.name,
                sale_type="credit",
                payment_method="cash",
                status="active",
                total_amount=Decimal("60.00"),
                paid_amount=Decimal("20.00"),
                balance_amount=Decimal("40.00"),
                credit_status="pending",
                created_by=user.id,
            ),
            Sale(
                sale_number="VTA-A-VOID",
                sale_date=date(2026, 7, 12),
                customer_id=customer_a.id,
                customer_name_snapshot=customer_a.name,
                sale_type="cash",
                payment_method="cash",
                status="voided",
                total_amount=Decimal("999.00"),
                paid_amount=Decimal("999.00"),
                balance_amount=Decimal("0.00"),
                credit_status="none",
                void_reason="Prueba",
                voided_at=voided_at,
                created_by=user.id,
            ),
            Sale(
                sale_number="VTA-A-OLD",
                sale_date=date(2026, 6, 30),
                customer_id=customer_a.id,
                customer_name_snapshot=customer_a.name,
                sale_type="cash",
                payment_method="cash",
                status="active",
                total_amount=Decimal("30.00"),
                paid_amount=Decimal("30.00"),
                balance_amount=Decimal("0.00"),
                credit_status="none",
                created_by=user.id,
            ),
            Sale(
                sale_number="VTA-B-CREDIT",
                sale_date=date(2026, 7, 11),
                customer_id=customer_b.id,
                customer_name_snapshot=customer_b.name,
                sale_type="credit",
                payment_method="cash",
                status="active",
                total_amount=Decimal("200.00"),
                paid_amount=Decimal("0.00"),
                balance_amount=Decimal("200.00"),
                credit_status="pending",
                created_by=user.id,
            ),
            Sale(
                sale_number="VTA-FINAL",
                sale_date=date(2026, 7, 12),
                customer_name_snapshot="Consumidor final",
                sale_type="cash",
                payment_method="cash",
                status="active",
                total_amount=Decimal("50.00"),
                paid_amount=Decimal("50.00"),
                balance_amount=Decimal("0.00"),
                credit_status="none",
                created_by=user.id,
            ),
        ]
        session.add_all(sales)
        session.flush()
        sale_by_number = {sale.sale_number: sale for sale in sales}
        session.add_all(
            [
                SaleItem(
                    sale_id=sale.id,
                    product_id=product.id,
                    product_sku_snapshot=product.sku,
                    product_name_snapshot=product.name,
                    quantity=Decimal("1.000"),
                    unit_price=sale.total_amount,
                    subtotal=sale.total_amount,
                )
                for sale in sales
            ]
        )
        credit_sale = sale_by_number["VTA-A-CREDIT"]
        session.add_all(
            [
                CreditPayment(
                    sale_id=credit_sale.id,
                    customer_id=customer_a.id,
                    payment_date=date(2026, 7, 12),
                    amount=Decimal("20.00"),
                    payment_method="cash",
                    status="active",
                    created_by=user.id,
                ),
                CreditPayment(
                    sale_id=credit_sale.id,
                    customer_id=customer_a.id,
                    payment_date=date(2026, 7, 12),
                    amount=Decimal("999.00"),
                    payment_method="cash",
                    status="voided",
                    void_reason="Prueba",
                    voided_at=voided_at,
                    created_by=user.id,
                ),
                OperatingExpense(
                    expense_number="GOP-62A",
                    expense_date=date(2026, 7, 12),
                    category_id=expense_category.id,
                    description="Gasto activo",
                    amount=Decimal("35.00"),
                    payment_status="paid",
                    payment_method="cash",
                    status="active",
                    created_by=user.id,
                ),
                OperatingExpense(
                    expense_number="GOP-62A-VOID",
                    expense_date=date(2026, 7, 12),
                    category_id=expense_category.id,
                    description="Gasto anulado",
                    amount=Decimal("999.00"),
                    payment_status="paid",
                    payment_method="cash",
                    status="voided",
                    void_reason="Prueba",
                    voided_at=voided_at,
                    created_by=user.id,
                ),
            ]
        )
        session.commit()
        return {
            "customer_a_id": customer_a.id,
            "customer_b_id": customer_b.id,
            "credit_sale_id": credit_sale.id,
        }


def _business_counts(testing_session_local) -> dict[str, int]:
    models = (Customer, Sale, SaleItem, CreditPayment, OperatingExpense)
    with testing_session_local() as session:
        return {
            model.__tablename__: int(
                session.scalar(select(func.count()).select_from(model)) or 0
            )
            for model in models
        }
