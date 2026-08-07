from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.models.credit_payment import CreditPayment
from app.models.inventory_movement import InventoryMovement
from app.models.sale import Sale
from app.models.simple_route import (
    SimpleRouteCollection,
    SimpleRouteExpense,
    SimpleRouteSale,
    SimpleRouteSheet,
    SimpleRouteTrip,
    SimpleRouteTripItem,
)
from app.services.excel_export import EXCEL_MEDIA_TYPE


SIMPLE_ROUTE_EXPORT_URLS = [
    "/api/v1/reports/exports/simple-routes-summary.xlsx",
    "/api/v1/reports/exports/simple-route-sales.xlsx",
    "/api/v1/reports/exports/simple-route-pending.xlsx",
    "/api/v1/reports/exports/simple-route-collections.xlsx",
    "/api/v1/reports/exports/simple-route-expenses.xlsx",
    "/api/v1/reports/exports/simple-route-settlements.xlsx",
]


@pytest.mark.parametrize("url", SIMPLE_ROUTE_EXPORT_URLS)
def test_admin_exports_valid_simple_route_xlsx_without_data(
    api_client,
    api_auth_headers,
    url: str,
) -> None:
    client, _ = api_client

    response = client.get(url, headers=api_auth_headers)

    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == EXCEL_MEDIA_TYPE
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert workbook.sheetnames
    assert any(
        "Sin " in str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    workbook.close()


@pytest.mark.parametrize(
    "url",
    [*SIMPLE_ROUTE_EXPORT_URLS, "/api/v1/reports/exports/simple-routes/999.xlsx"],
)
def test_operator_cannot_export_simple_routes(
    api_client,
    api_operator_auth_headers,
    url: str,
) -> None:
    client, _ = api_client

    response = client.get(url, headers=api_operator_auth_headers)

    assert response.status_code == 403
    assert "permisos" in response.json()["detail"].lower()


def test_summary_uses_only_active_simple_routes_for_main_totals(
    api_client,
    api_auth_headers,
) -> None:
    client, _ = api_client
    sheet_id = _seed_complete_route(client, api_auth_headers)
    cancelled_id = _create_sheet(
        client,
        api_auth_headers,
        seller="Vendedor cancelado",
        route_label="Ruta cancelada",
        day=2,
    )
    _add_sale(
        client,
        api_auth_headers,
        cancelled_id,
        sale_type="cash",
        quantity="10",
        unit_price="10",
    )
    cancelled = client.post(
        f"/api/v1/simple-routes/{cancelled_id}/cancel",
        headers=api_auth_headers,
        json={"reason": "Ruta no realizada"},
    )
    assert cancelled.status_code == 200

    response = client.get(
        "/api/v1/reports/exports/simple-routes-summary.xlsx"
        "?start_date=2026-08-01&end_date=2026-08-31",
        headers=api_auth_headers,
    )

    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    summary = _summary_values(workbook["Resumen"])
    assert summary["Total rutas"] == 2
    assert summary["Rutas liquidadas"] == 1
    assert summary["Rutas canceladas"] == 1
    assert summary["Ventas de ruta"] == 40
    assert summary["Fiado de rutas"] == 17
    assert summary["Abonos de rutas"] == 4
    assert summary["Gastos de rutas"] == 2
    assert summary["Efectivo esperado"] == 25
    assert summary["Efectivo entregado"] == 24
    assert summary["Diferencia total"] == -1
    route_rows = list(workbook["Rutas"].iter_rows(min_row=2, values_only=True))
    assert {row[1] for row in route_rows} == {sheet_id, cancelled_id}
    assert any(row[5] == "Cancelada" for row in route_rows)
    workbook.close()


def test_detail_exports_all_simple_route_sections_and_sanitizes_formulas(
    api_client,
    api_auth_headers,
) -> None:
    client, _ = api_client
    sheet_id = _seed_complete_route(client, api_auth_headers, notes="=2+2")

    response = client.get(
        f"/api/v1/reports/exports/simple-routes/{sheet_id}.xlsx",
        headers=api_auth_headers,
    )

    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content), data_only=False)
    assert workbook.sheetnames == [
        "Encabezado",
        "Viajes",
        "Ventas",
        "Abonos",
        "Gastos",
        "Cuadre",
    ]
    header = _summary_values(workbook["Encabezado"])
    assert header["Ruta ID"] == sheet_id
    assert header["Estado"] == "Liquidada"
    assert header["Observaciones"] == "'=2+2"
    assert workbook["Viajes"].max_row == 3
    assert workbook["Ventas"].max_row == 4
    assert workbook["Abonos"]["E2"].value == 4
    assert workbook["Gastos"]["E2"].value == 2
    assert workbook["Cuadre"]["G2"].value == 0
    assert workbook["Cuadre"]["J2"].value == -1
    workbook.close()


def test_detail_for_missing_simple_route_returns_404(
    api_client,
    api_auth_headers,
) -> None:
    client, _ = api_client

    response = client.get(
        "/api/v1/reports/exports/simple-routes/999.xlsx",
        headers=api_auth_headers,
    )

    assert response.status_code == 404
    assert "no existe" in response.json()["detail"]


def test_sales_pending_collections_expenses_and_settlements_are_detailed(
    api_client,
    api_auth_headers,
) -> None:
    client, _ = api_client
    sheet_id = _seed_complete_route(client, api_auth_headers)

    sales_response = client.get(
        "/api/v1/reports/exports/simple-route-sales.xlsx"
        "?start_date=2026-08-01&end_date=2026-08-01&seller=George"
        "&route_label=Centro&sale_type=all",
        headers=api_auth_headers,
    )
    pending_response = client.get(
        "/api/v1/reports/exports/simple-route-pending.xlsx"
        "?start_date=2026-08-01&end_date=2026-08-01",
        headers=api_auth_headers,
    )
    collections_response = client.get(
        "/api/v1/reports/exports/simple-route-collections.xlsx"
        "?start_date=2026-08-01&end_date=2026-08-01",
        headers=api_auth_headers,
    )
    expenses_response = client.get(
        "/api/v1/reports/exports/simple-route-expenses.xlsx"
        "?start_date=2026-08-01&end_date=2026-08-01",
        headers=api_auth_headers,
    )
    settlements_response = client.get(
        "/api/v1/reports/exports/simple-route-settlements.xlsx"
        "?start_date=2026-08-01&end_date=2026-08-01&status=settled",
        headers=api_auth_headers,
    )

    for response in (
        sales_response,
        pending_response,
        collections_response,
        expenses_response,
        settlements_response,
    ):
        assert response.status_code == 200, response.text

    sales_workbook = load_workbook(BytesIO(sales_response.content), data_only=True)
    summary = _summary_values(sales_workbook["Resumen ventas ruta"])
    assert summary["Total ventas ruta"] == 40
    assert summary["Contado"] == 20
    assert summary["Fiado"] == 12
    assert summary["Parcial"] == 8
    assert summary["Cobrado"] == 27
    assert summary["Saldo fiado"] == 13
    assert summary["Cantidad de ventas"] == 3
    assert {
        row[6]
        for row in sales_workbook["Detalle ventas"].iter_rows(
            min_row=2, values_only=True
        )
    } == {"Contado", "Fiado", "Parcial"}
    sales_workbook.close()

    pending_workbook = load_workbook(BytesIO(pending_response.content))
    pending_rows = list(
        pending_workbook["Fiados pendientes"].iter_rows(
            min_row=2, values_only=True
        )
    )
    assert {row[9] for row in pending_rows} == {8, 5}
    pending_workbook.close()

    collections_workbook = load_workbook(BytesIO(collections_response.content))
    assert collections_workbook["Abonos de ruta"]["H2"].value == 4
    collections_workbook.close()

    expenses_workbook = load_workbook(BytesIO(expenses_response.content))
    assert expenses_workbook["Gastos de ruta"]["H2"].value == 2
    expenses_workbook.close()

    settlements_workbook = load_workbook(BytesIO(settlements_response.content))
    settlement_row = next(
        settlements_workbook["Cuadres por ruta"].iter_rows(
            min_row=2, values_only=True
        )
    )
    assert settlement_row[1] == sheet_id
    assert settlement_row[8:11] == (25, 24, -1)
    product_row = next(
        settlements_workbook["Diferencias de producto"].iter_rows(
            min_row=2, values_only=True
        )
    )
    assert product_row[5:10] == (5, 1, 4, 4, 0)
    settlements_workbook.close()


def test_simple_route_export_filters_and_validation(
    api_client,
    api_auth_headers,
) -> None:
    client, _ = api_client
    _seed_complete_route(client, api_auth_headers)
    _create_sheet(
        client,
        api_auth_headers,
        seller="Ana",
        route_label="Sur",
        day=5,
    )

    filtered = client.get(
        "/api/v1/reports/exports/simple-routes-summary.xlsx"
        "?start_date=2026-08-01&end_date=2026-08-01&seller=geo&route_label=cent",
        headers=api_auth_headers,
    )
    inverted = client.get(
        "/api/v1/reports/exports/simple-route-sales.xlsx"
        "?start_date=2026-08-31&end_date=2026-08-01",
        headers=api_auth_headers,
    )
    invalid_type = client.get(
        "/api/v1/reports/exports/simple-route-sales.xlsx?sale_type=official",
        headers=api_auth_headers,
    )

    workbook = load_workbook(BytesIO(filtered.content), data_only=True)
    assert _summary_values(workbook["Resumen"])["Total rutas"] == 1
    workbook.close()
    assert inverted.status_code == 400
    assert invalid_type.status_code == 400


def test_simple_route_exports_do_not_modify_or_mix_official_modules(
    api_client,
    api_auth_headers,
) -> None:
    client, session_factory = api_client
    sheet_id = _seed_complete_route(client, api_auth_headers)
    before = _record_counts(session_factory)

    urls = [
        *SIMPLE_ROUTE_EXPORT_URLS,
        f"/api/v1/reports/exports/simple-routes/{sheet_id}.xlsx",
    ]
    for url in urls:
        response = client.get(url, headers=api_auth_headers)
        assert response.status_code == 200, response.text

    assert _record_counts(session_factory) == before


def _seed_complete_route(client, headers, notes: str | None = None) -> int:
    sheet_id = _create_sheet(client, headers, notes=notes)
    trip_id = client.post(
        f"/api/v1/simple-routes/{sheet_id}/trips",
        headers=headers,
        json={"notes": "Primer viaje"},
    ).json()["id"]
    for movement, quantity in (("out", "5"), ("return", "1")):
        response = client.post(
            f"/api/v1/simple-routes/{sheet_id}/trips/{trip_id}/items",
            headers=headers,
            json={
                "movement_type": movement,
                "item_label": "Bidón lleno",
                "item_category": "full_jug",
                "quantity": quantity,
            },
        )
        assert response.status_code == 201, response.text

    _add_sale(client, headers, sheet_id, "cash", quantity="2", unit_price="10")
    credit = _add_sale(
        client,
        headers,
        sheet_id,
        "credit",
        quantity="1",
        unit_price="12",
    ).json()
    _add_sale(
        client,
        headers,
        sheet_id,
        "partial",
        paid="3",
        quantity="1",
        unit_price="8",
    )
    collection = client.post(
        f"/api/v1/simple-routes/{sheet_id}/collections",
        headers=headers,
        json={
            "simple_route_sale_id": credit["id"],
            "amount": "4",
            "collection_date": "2026-08-01",
            "notes": "Abono parcial",
        },
    )
    expense = client.post(
        f"/api/v1/simple-routes/{sheet_id}/expenses",
        headers=headers,
        json={
            "description": "Combustible",
            "category": "fuel",
            "amount": "2",
            "expense_date": "2026-08-01",
        },
    )
    assert collection.status_code == 201, collection.text
    assert expense.status_code == 201, expense.text
    settled = client.post(
        f"/api/v1/simple-routes/{sheet_id}/settle",
        headers=headers,
        json={"cash_delivered": "24"},
    )
    assert settled.status_code == 200, settled.text
    return sheet_id


def _create_sheet(
    client,
    headers,
    *,
    seller: str = "George",
    route_label: str = "Centro",
    day: int = 1,
    notes: str | None = None,
) -> int:
    response = client.post(
        "/api/v1/simple-routes",
        headers=headers,
        json={
            "route_date": f"2026-08-{day:02d}",
            "seller_name": seller,
            "route_label": route_label,
            "notes": notes,
        },
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


def _add_sale(
    client,
    headers,
    sheet_id: int,
    sale_type: str,
    paid: str | None = None,
    *,
    quantity: str,
    unit_price: str,
):
    payload = {
        "customer_name": "Cliente de ruta",
        "sale_type": sale_type,
        "sale_date": "2026-08-01",
        "notes": "Venta simple",
        "items": [
            {
                "product_name": "Bidón lleno",
                "item_category": "full_jug",
                "quantity": quantity,
                "unit_price": unit_price,
            }
        ],
    }
    if paid is not None:
        payload["paid_amount"] = paid
    response = client.post(
        f"/api/v1/simple-routes/{sheet_id}/sales",
        headers=headers,
        json=payload,
    )
    assert response.status_code == 201, response.text
    return response


def _summary_values(sheet) -> dict[str, object]:
    return {
        str(label): value
        for label, value, *_ in sheet.iter_rows(min_row=2, values_only=True)
    }


def _record_counts(session_factory) -> dict[str, int]:
    models = (
        SimpleRouteSheet,
        SimpleRouteTrip,
        SimpleRouteTripItem,
        SimpleRouteSale,
        SimpleRouteCollection,
        SimpleRouteExpense,
        Sale,
        CreditPayment,
        InventoryMovement,
    )
    with session_factory() as session:
        return {
            model.__tablename__: int(
                session.scalar(select(func.count()).select_from(model)) or 0
            )
            for model in models
        }
