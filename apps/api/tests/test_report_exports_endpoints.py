from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.models.cash_session import CashSession
from app.models.credit_payment import CreditPayment
from app.models.customer import Customer
from app.models.employee import Employee
from app.models.inventory_movement import InventoryMovement
from app.models.inventory_stock import InventoryStock
from app.models.operating_expense import OperatingExpense
from app.models.payroll_run import PayrollRun
from app.models.production_batch import ProductionBatch
from app.models.production_inventory_application import (
    ProductionInventoryApplication,
)
from app.models.production_material_consumption import (
    ProductionMaterialConsumption,
)
from app.models.purchase import Purchase
from app.models.route_cash_receipt import RouteCashReceipt
from app.models.route_closure import RouteClosure
from app.models.route_load import RouteLoad
from app.models.route_return import RouteReturn
from app.models.sale import Sale
from app.services.excel_export import EXCEL_MEDIA_TYPE
from tests.test_management_dashboard_endpoints import _seed_management_data

EXPORT_URLS = [
    "/api/v1/reports/exports/management-dashboard.xlsx",
    "/api/v1/reports/exports/monthly-summary.xlsx?year=2026&month=7",
    "/api/v1/reports/exports/sales.xlsx",
    "/api/v1/reports/exports/receivables.xlsx",
    "/api/v1/reports/exports/operating-expenses.xlsx",
    "/api/v1/reports/exports/purchases.xlsx",
    "/api/v1/reports/exports/inventory.xlsx",
    "/api/v1/reports/exports/production.xlsx",
    "/api/v1/reports/exports/routes.xlsx",
    "/api/v1/reports/exports/payroll.xlsx",
]

DATED_EXPORT_URLS = [
    "/api/v1/reports/exports/management-dashboard.xlsx",
    "/api/v1/reports/exports/sales.xlsx",
    "/api/v1/reports/exports/receivables.xlsx",
    "/api/v1/reports/exports/operating-expenses.xlsx",
    "/api/v1/reports/exports/purchases.xlsx",
    "/api/v1/reports/exports/production.xlsx",
    "/api/v1/reports/exports/routes.xlsx",
    "/api/v1/reports/exports/payroll.xlsx",
]

BUSINESS_MODELS = [
    CashSession,
    Sale,
    CreditPayment,
    OperatingExpense,
    Purchase,
    InventoryStock,
    InventoryMovement,
    ProductionBatch,
    ProductionInventoryApplication,
    ProductionMaterialConsumption,
    RouteLoad,
    RouteReturn,
    RouteClosure,
    RouteCashReceipt,
    Employee,
    PayrollRun,
]


@pytest.mark.parametrize("url", EXPORT_URLS)
def test_report_export_requires_authentication(api_client, url: str) -> None:
    client, _ = api_client

    response = client.get(url)

    assert response.status_code == 401


@pytest.mark.parametrize("url", EXPORT_URLS)
def test_report_export_returns_valid_xlsx_without_data(
    api_client,
    api_auth_headers,
    url: str,
) -> None:
    client, _ = api_client

    response = client.get(url, headers=api_auth_headers)

    assert response.status_code == 200
    assert response.headers["content-type"] == EXCEL_MEDIA_TYPE
    assert ".xlsx" in response.headers["content-disposition"]

    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert workbook.sheetnames
    for sheet in workbook.worksheets:
        assert sheet.max_column >= 1
        assert sheet.cell(row=1, column=1).value
    workbook.close()


def test_management_export_contains_expected_sheets(
    api_client,
    api_auth_headers,
) -> None:
    client, _ = api_client

    response = client.get(EXPORT_URLS[0], headers=api_auth_headers)

    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert {
        "Resumen",
        "Caja",
        "Ventas",
        "Cartera",
        "Gastos operativos",
        "Compras",
        "Producción",
        "Inventario",
        "Rutas",
        "Nómina",
        "Alertas",
        "Series diarias",
    }.issubset(workbook.sheetnames)
    workbook.close()


def test_monthly_export_contains_expected_sheets(
    api_client,
    api_auth_headers,
) -> None:
    client, _ = api_client

    response = client.get(EXPORT_URLS[1], headers=api_auth_headers)

    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert {
        "Resumen mensual",
        "Caja",
        "Ventas",
        "Cartera",
        "Producción",
        "Rutas",
        "Recibos",
        "Alertas",
        "Días del mes",
    }.issubset(workbook.sheetnames)
    workbook.close()


@pytest.mark.parametrize("url", DATED_EXPORT_URLS)
def test_report_export_rejects_inverted_date_range(
    api_client,
    api_auth_headers,
    url: str,
) -> None:
    client, _ = api_client
    separator = "&" if "?" in url else "?"

    response = client.get(
        f"{url}{separator}date_from=2026-07-31&date_to=2026-07-01",
        headers=api_auth_headers,
    )

    assert response.status_code == 400


def test_monthly_export_requires_year_and_month(
    api_client,
    api_auth_headers,
) -> None:
    client, _ = api_client

    missing_period = client.get(
        "/api/v1/reports/exports/monthly-summary.xlsx",
        headers=api_auth_headers,
    )
    invalid_month = client.get(
        "/api/v1/reports/exports/monthly-summary.xlsx?year=2026&month=13",
        headers=api_auth_headers,
    )

    assert missing_period.status_code == 422
    assert invalid_month.status_code == 422


def test_all_report_exports_are_read_only(
    api_client,
    api_auth_headers,
) -> None:
    client, testing_session_local = api_client
    before = _business_record_counts(testing_session_local)

    for url in EXPORT_URLS:
        response = client.get(url, headers=api_auth_headers)
        assert response.status_code == 200

    assert _business_record_counts(testing_session_local) == before


def test_report_exports_open_with_registered_business_data(
    api_client,
    api_auth_headers,
) -> None:
    client, testing_session_local = api_client
    _seed_management_data(testing_session_local)
    filtered_urls = [
        (
            f"{url}{'&' if '?' in url else '?'}"
            "date_from=2026-07-01&date_to=2026-07-31"
        )
        if "monthly-summary" not in url and "inventory" not in url
        else url
        for url in EXPORT_URLS
    ]

    for url in filtered_urls:
        response = client.get(url, headers=api_auth_headers)
        assert response.status_code == 200
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        assert workbook.sheetnames
        workbook.close()


def test_receivables_export_applies_customer_filter_to_summary(
    api_client,
    api_auth_headers,
) -> None:
    client, testing_session_local = api_client
    _seed_management_data(testing_session_local)
    with testing_session_local() as session:
        customer_id = session.scalar(
            select(Customer.id).where(Customer.name == "Cliente crédito")
        )
    assert customer_id is not None

    response = client.get(
        (
            "/api/v1/reports/exports/receivables.xlsx"
            f"?customer_id={customer_id}"
        ),
        headers=api_auth_headers,
    )

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    summary = workbook["Resumen"]
    assert summary["B2"].value == 1
    assert summary["B3"].value == 60
    assert summary["B5"].value == 40
    workbook.close()


def _business_record_counts(testing_session_local) -> dict[str, int]:
    with testing_session_local() as session:
        return {
            model.__tablename__: int(
                session.scalar(select(func.count()).select_from(model)) or 0
            )
            for model in BUSINESS_MODELS
        }
