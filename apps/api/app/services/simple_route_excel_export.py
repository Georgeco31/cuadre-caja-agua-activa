from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from sqlalchemy.orm import Session

from app.models.simple_route import SimpleRouteSale, SimpleRouteSheet
from app.repositories import simple_routes as simple_route_repository
from app.services import simple_routes as simple_route_service
from app.services.excel_export import add_sheet, create_workbook


ZERO_MONEY = Decimal("0.00")
ZERO_QUANTITY = Decimal("0.000")
EMPTY_PERIOD_MESSAGE = "Sin Rutas Simples en el período"

STATUS_LABELS = {
    "draft": "Borrador",
    "in_progress": "En progreso",
    "settled": "Liquidada",
    "cancelled": "Cancelada",
}
SALE_TYPE_LABELS = {
    "cash": "Contado",
    "credit": "Fiado",
    "partial": "Parcial",
}
MOVEMENT_LABELS = {"out": "Sale", "return": "Regresa"}


def build_simple_routes_summary_workbook(
    db: Session,
    *,
    start_date: date | None = None,
    end_date: date | None = None,
    seller: str | None = None,
    route_label: str | None = None,
    status: str | None = None,
) -> Workbook:
    sheets = _list_filtered_sheets(
        db,
        start_date=start_date,
        end_date=end_date,
        seller=seller,
        route_label=route_label,
        status=status,
        period_field="route",
    )
    workbook = create_workbook()
    settlements = {
        sheet.id: simple_route_service.calculate_settlement(sheet) for sheet in sheets
    }
    active_sheets = [sheet for sheet in sheets if sheet.status != "cancelled"]
    active_settlements = [settlements[sheet.id] for sheet in active_sheets]
    status_counts = {
        route_status: sum(sheet.status == route_status for sheet in sheets)
        for route_status in STATUS_LABELS
    }

    summary_rows: list[tuple[Any, Any]] = [
        ("Fecha desde", start_date or "Todas"),
        ("Fecha hasta", end_date or "Todas"),
        ("Vendedor", seller or "Todos"),
        ("Ruta / zona", route_label or "Todas"),
        ("Estado", STATUS_LABELS.get(status, "Todos")),
        ("Total rutas", len(sheets)),
        ("Rutas borrador", status_counts["draft"]),
        ("Rutas en progreso", status_counts["in_progress"]),
        ("Rutas liquidadas", status_counts["settled"]),
        ("Rutas canceladas", status_counts["cancelled"]),
        (
            "Ventas de ruta",
            _sum_money(item.route_sales_total for item in active_settlements),
        ),
        (
            "Fiado de rutas",
            _sum_money(item.route_credit_total for item in active_settlements),
        ),
        (
            "Abonos de rutas",
            _sum_money(item.collections_total for item in active_settlements),
        ),
        (
            "Gastos de rutas",
            _sum_money(item.expenses_total for item in active_settlements),
        ),
        (
            "Efectivo esperado",
            _sum_money(item.expected_cash for item in active_settlements),
        ),
        (
            "Efectivo entregado",
            _sum_money(
                item.cash_delivered or ZERO_MONEY for item in active_settlements
            ),
        ),
        (
            "Diferencia total",
            _sum_money(
                item.cash_difference or ZERO_MONEY for item in active_settlements
            ),
        ),
        ("Fecha de generación", datetime.now().replace(microsecond=0)),
    ]
    if not sheets:
        summary_rows.append(("Resultado", EMPTY_PERIOD_MESSAGE))
    summary_sheet = add_sheet(
        workbook,
        "Resumen",
        ("Indicador", "Valor"),
        summary_rows,
    )
    for row_number in range(12, 19):
        summary_sheet.cell(row=row_number, column=2).number_format = '"$"#,##0.00'

    add_sheet(
        workbook,
        "Rutas",
        (
            "Fecha",
            "Ruta ID",
            "Vendedor",
            "Vehículo",
            "Ruta / zona",
            "Estado",
            "Ventas ruta",
            "Fiado",
            "Abonos",
            "Gastos",
            "Efectivo esperado",
            "Efectivo entregado",
            "Diferencia",
            "Observación",
            "Motivo cancelación",
        ),
        _rows_or_message(
            [
                (
                    sheet.route_date,
                    sheet.id,
                    sheet.seller_name,
                    sheet.vehicle_name,
                    sheet.route_label,
                    STATUS_LABELS[sheet.status],
                    settlements[sheet.id].route_sales_total,
                    settlements[sheet.id].route_credit_total,
                    settlements[sheet.id].collections_total,
                    settlements[sheet.id].expenses_total,
                    settlements[sheet.id].expected_cash,
                    settlements[sheet.id].cash_delivered,
                    settlements[sheet.id].cash_difference,
                    sheet.notes,
                    sheet.cancellation_reason,
                )
                for sheet in sheets
            ],
            15,
            EMPTY_PERIOD_MESSAGE,
        ),
        money_columns=(
            "Ventas ruta",
            "Fiado",
            "Abonos",
            "Gastos",
            "Efectivo esperado",
            "Efectivo entregado",
            "Diferencia",
        ),
        date_columns=("Fecha",),
    )
    _add_rankings_sheet(workbook, active_sheets, settlements)
    return workbook


def build_simple_route_detail_workbook(
    db: Session,
    route_sheet_id: int,
) -> Workbook:
    sheet = simple_route_service.get_sheet(db, route_sheet_id)
    settlement = simple_route_service.calculate_settlement(sheet)
    workbook = create_workbook()

    header_sheet = add_sheet(
        workbook,
        "Encabezado",
        ("Campo", "Valor"),
        (
            ("Ruta ID", sheet.id),
            ("Fecha", sheet.route_date),
            ("Vendedor", sheet.seller_name),
            ("Empleado relacionado", sheet.employee_name),
            ("Vehículo", sheet.vehicle_name),
            ("Ruta / zona", sheet.route_label),
            ("Estado", STATUS_LABELS[sheet.status]),
            ("Observaciones", sheet.notes),
            ("Motivo cancelación", sheet.cancellation_reason),
            ("Dinero entregado", settlement.cash_delivered),
            ("Efectivo esperado", settlement.expected_cash),
            ("Diferencia", settlement.cash_difference),
        ),
    )
    for row_number in range(11, 14):
        header_sheet.cell(row=row_number, column=2).number_format = '"$"#,##0.00'

    trip_rows = [
        (
            trip.trip_number,
            MOVEMENT_LABELS[item.movement_type],
            item.item_label,
            item.item_category,
            item.quantity,
            item.notes or trip.notes,
        )
        for trip in sheet.trips
        for item in trip.items
    ]
    add_sheet(
        workbook,
        "Viajes",
        (
            "Número de viaje",
            "Movimiento",
            "Producto / etiqueta",
            "Categoría",
            "Cantidad",
            "Observación",
        ),
        _rows_or_message(trip_rows, 6, "Sin viajes registrados"),
        quantity_columns=("Cantidad",),
    )

    add_sheet(
        workbook,
        "Ventas",
        (
            "Fecha",
            "Venta ID",
            "Cliente",
            "Tipo",
            "Producto",
            "Categoría",
            "Cantidad",
            "Precio unitario",
            "Subtotal",
            "Total venta",
            "Cobrado",
            "Saldo",
            "Estado venta",
            "Observación",
        ),
        _rows_or_message(_sale_rows(sheet), 14, "Sin ventas de ruta"),
        money_columns=(
            "Precio unitario",
            "Subtotal",
            "Total venta",
            "Cobrado",
            "Saldo",
        ),
        quantity_columns=("Cantidad",),
        date_columns=("Fecha",),
    )

    add_sheet(
        workbook,
        "Abonos",
        ("Fecha", "Abono ID", "Venta ID", "Cliente", "Monto", "Observación"),
        _rows_or_message(
            [
                (
                    item.collection_date,
                    item.id,
                    item.simple_route_sale_id,
                    item.customer_name_snapshot,
                    item.amount,
                    item.notes,
                )
                for item in sheet.collections
            ],
            6,
            "Sin abonos de ruta",
        ),
        money_columns=("Monto",),
        date_columns=("Fecha",),
    )

    add_sheet(
        workbook,
        "Gastos",
        ("Fecha", "Gasto ID", "Categoría", "Descripción", "Monto", "Observación"),
        _rows_or_message(
            [
                (
                    item.expense_date,
                    item.id,
                    item.category,
                    item.description,
                    item.amount,
                    item.notes,
                )
                for item in sheet.expenses
            ],
            6,
            "Sin gastos de ruta",
        ),
        money_columns=("Monto",),
        date_columns=("Fecha",),
    )

    settlement_rows = [
        (
            item.item_label,
            item.item_category,
            item.total_out,
            item.total_returned,
            item.physical_sold,
            item.registered_sold,
            item.physical_difference,
            settlement.expected_cash,
            settlement.cash_delivered,
            settlement.cash_difference,
        )
        for item in settlement.products
    ]
    if not settlement_rows:
        settlement_rows = [
            (
                "Sin movimientos de producto",
                None,
                ZERO_QUANTITY,
                ZERO_QUANTITY,
                ZERO_QUANTITY,
                ZERO_QUANTITY,
                ZERO_QUANTITY,
                settlement.expected_cash,
                settlement.cash_delivered,
                settlement.cash_difference,
            )
        ]
    add_sheet(
        workbook,
        "Cuadre",
        (
            "Producto",
            "Categoría",
            "Total salido",
            "Total regresado",
            "Vendido físico",
            "Vendido registrado",
            "Diferencia producto",
            "Efectivo esperado",
            "Efectivo entregado",
            "Diferencia efectivo",
        ),
        settlement_rows,
        money_columns=(
            "Efectivo esperado",
            "Efectivo entregado",
            "Diferencia efectivo",
        ),
        quantity_columns=(
            "Total salido",
            "Total regresado",
            "Vendido físico",
            "Vendido registrado",
            "Diferencia producto",
        ),
    )
    return workbook


def build_simple_route_sales_workbook(
    db: Session,
    *,
    start_date: date | None = None,
    end_date: date | None = None,
    sale_type: str | None = None,
    seller: str | None = None,
    route_label: str | None = None,
) -> Workbook:
    _validate_sale_type(sale_type)
    sheets = _list_filtered_sheets(
        db,
        start_date=start_date,
        end_date=end_date,
        seller=seller,
        route_label=route_label,
        period_field="events",
    )
    sales = [
        sale
        for sheet in sheets
        for sale in sheet.sales
        if sale.status == "active"
        and _date_in_period(sale.sale_date, start_date, end_date)
        and (sale_type in (None, "all") or sale.sale_type == sale_type)
    ]
    workbook = create_workbook()
    sale_totals = {
        route_sale_type: _sum_money(
            sale.total_amount for sale in sales if sale.sale_type == route_sale_type
        )
        for route_sale_type in SALE_TYPE_LABELS
    }
    collected = _sum_money(
        sale.total_amount - sale.remaining_amount for sale in sales
    )
    pending = _sum_money(sale.remaining_amount for sale in sales)
    summary_rows: list[tuple[Any, Any]] = [
        ("Fecha desde", start_date or "Todas"),
        ("Fecha hasta", end_date or "Todas"),
        ("Total ventas ruta", _sum_money(sale.total_amount for sale in sales)),
        ("Contado", sale_totals["cash"]),
        ("Fiado", sale_totals["credit"]),
        ("Parcial", sale_totals["partial"]),
        ("Cobrado", collected),
        ("Saldo fiado", pending),
        ("Cantidad de ventas", len(sales)),
    ]
    if not sales:
        summary_rows.append(("Resultado", "Sin ventas de Rutas Simples en el período"))
    summary_sheet = add_sheet(
        workbook,
        "Resumen ventas ruta",
        ("Indicador", "Valor"),
        summary_rows,
    )
    for row_number in range(4, 10):
        summary_sheet.cell(row=row_number, column=2).number_format = '"$"#,##0.00'
    add_sheet(
        workbook,
        "Detalle ventas",
        (
            "Fecha",
            "Ruta ID",
            "Vendedor",
            "Ruta / zona",
            "Venta ID",
            "Cliente",
            "Tipo",
            "Producto",
            "Categoría",
            "Cantidad",
            "Precio",
            "Subtotal",
            "Total venta",
            "Cobrado",
            "Saldo",
            "Observación",
        ),
        _rows_or_message(
            [
                (
                    sale.sale_date,
                    sale.route_sheet.id,
                    sale.route_sheet.seller_name,
                    sale.route_sheet.route_label,
                    sale.id,
                    sale.customer_name_snapshot,
                    SALE_TYPE_LABELS[sale.sale_type],
                    item.product_name_snapshot,
                    item.item_category,
                    item.quantity,
                    item.unit_price,
                    item.subtotal,
                    sale.total_amount,
                    sale.total_amount - sale.remaining_amount,
                    sale.remaining_amount,
                    sale.notes,
                )
                for sale in sales
                for item in sale.items
            ],
            16,
            "Sin ventas de Rutas Simples en el período",
        ),
        money_columns=("Precio", "Subtotal", "Total venta", "Cobrado", "Saldo"),
        quantity_columns=("Cantidad",),
        date_columns=("Fecha",),
    )
    return workbook


def build_simple_route_pending_workbook(
    db: Session,
    *,
    start_date: date | None = None,
    end_date: date | None = None,
    seller: str | None = None,
    route_label: str | None = None,
) -> Workbook:
    sheets = _list_filtered_sheets(
        db,
        start_date=start_date,
        end_date=end_date,
        seller=seller,
        route_label=route_label,
        period_field="events",
    )
    sales = [
        sale
        for sheet in sheets
        for sale in sheet.sales
        if sale.status == "active"
        and sale.remaining_amount > ZERO_MONEY
        and _date_in_period(sale.sale_date, start_date, end_date)
    ]
    workbook = create_workbook()
    add_sheet(
        workbook,
        "Fiados pendientes",
        (
            "Fecha",
            "Ruta ID",
            "Vendedor",
            "Ruta / zona",
            "Venta ID",
            "Cliente",
            "Tipo",
            "Total venta",
            "Cobrado",
            "Saldo pendiente",
            "Observación",
        ),
        _rows_or_message(
            [
                (
                    sale.sale_date,
                    sale.route_sheet.id,
                    sale.route_sheet.seller_name,
                    sale.route_sheet.route_label,
                    sale.id,
                    sale.customer_name_snapshot,
                    SALE_TYPE_LABELS[sale.sale_type],
                    sale.total_amount,
                    sale.total_amount - sale.remaining_amount,
                    sale.remaining_amount,
                    sale.notes,
                )
                for sale in sales
            ],
            11,
            "Sin fiados pendientes de Rutas Simples en el período",
        ),
        money_columns=("Total venta", "Cobrado", "Saldo pendiente"),
        date_columns=("Fecha",),
    )
    return workbook


def build_simple_route_collections_workbook(
    db: Session,
    *,
    start_date: date | None = None,
    end_date: date | None = None,
    seller: str | None = None,
    route_label: str | None = None,
) -> Workbook:
    sheets = _list_filtered_sheets(
        db,
        start_date=start_date,
        end_date=end_date,
        seller=seller,
        route_label=route_label,
        period_field="events",
    )
    collections = [
        item
        for sheet in sheets
        for item in sheet.collections
        if _date_in_period(item.collection_date, start_date, end_date)
    ]
    workbook = create_workbook()
    add_sheet(
        workbook,
        "Abonos de ruta",
        (
            "Fecha",
            "Ruta ID",
            "Vendedor",
            "Ruta / zona",
            "Abono ID",
            "Venta ID",
            "Cliente",
            "Monto",
            "Observación",
        ),
        _rows_or_message(
            [
                (
                    item.collection_date,
                    item.route_sheet.id,
                    item.route_sheet.seller_name,
                    item.route_sheet.route_label,
                    item.id,
                    item.simple_route_sale_id,
                    item.customer_name_snapshot,
                    item.amount,
                    item.notes,
                )
                for item in collections
            ],
            9,
            "Sin abonos de Rutas Simples en el período",
        ),
        money_columns=("Monto",),
        date_columns=("Fecha",),
    )
    return workbook


def build_simple_route_expenses_workbook(
    db: Session,
    *,
    start_date: date | None = None,
    end_date: date | None = None,
    seller: str | None = None,
    route_label: str | None = None,
) -> Workbook:
    sheets = _list_filtered_sheets(
        db,
        start_date=start_date,
        end_date=end_date,
        seller=seller,
        route_label=route_label,
        period_field="events",
    )
    expenses = [
        item
        for sheet in sheets
        for item in sheet.expenses
        if _date_in_period(item.expense_date, start_date, end_date)
    ]
    workbook = create_workbook()
    add_sheet(
        workbook,
        "Gastos de ruta",
        (
            "Fecha",
            "Ruta ID",
            "Vendedor",
            "Ruta / zona",
            "Gasto ID",
            "Categoría",
            "Descripción",
            "Monto",
            "Observación",
        ),
        _rows_or_message(
            [
                (
                    item.expense_date,
                    item.route_sheet.id,
                    item.route_sheet.seller_name,
                    item.route_sheet.route_label,
                    item.id,
                    item.category,
                    item.description,
                    item.amount,
                    item.notes,
                )
                for item in expenses
            ],
            9,
            "Sin gastos de Rutas Simples en el período",
        ),
        money_columns=("Monto",),
        date_columns=("Fecha",),
    )
    return workbook


def build_simple_route_settlements_workbook(
    db: Session,
    *,
    start_date: date | None = None,
    end_date: date | None = None,
    seller: str | None = None,
    route_label: str | None = None,
    status: str | None = None,
) -> Workbook:
    sheets = _list_filtered_sheets(
        db,
        start_date=start_date,
        end_date=end_date,
        seller=seller,
        route_label=route_label,
        status=status,
        period_field="route",
    )
    settlements = {
        sheet.id: simple_route_service.calculate_settlement(sheet) for sheet in sheets
    }
    workbook = create_workbook()
    add_sheet(
        workbook,
        "Cuadres por ruta",
        (
            "Fecha",
            "Ruta ID",
            "Vendedor",
            "Ruta / zona",
            "Ventas ruta",
            "Fiado",
            "Abonos",
            "Gastos",
            "Efectivo esperado",
            "Efectivo entregado",
            "Diferencia efectivo",
            "Estado",
            "Estado cuadre",
            "Observación",
            "Motivo cancelación",
        ),
        _rows_or_message(
            [
                (
                    sheet.route_date,
                    sheet.id,
                    sheet.seller_name,
                    sheet.route_label,
                    settlements[sheet.id].route_sales_total,
                    settlements[sheet.id].route_credit_total,
                    settlements[sheet.id].collections_total,
                    settlements[sheet.id].expenses_total,
                    settlements[sheet.id].expected_cash,
                    settlements[sheet.id].cash_delivered,
                    settlements[sheet.id].cash_difference,
                    STATUS_LABELS[sheet.status],
                    settlements[sheet.id].status,
                    sheet.notes,
                    sheet.cancellation_reason,
                )
                for sheet in sheets
            ],
            15,
            EMPTY_PERIOD_MESSAGE,
        ),
        money_columns=(
            "Ventas ruta",
            "Fiado",
            "Abonos",
            "Gastos",
            "Efectivo esperado",
            "Efectivo entregado",
            "Diferencia efectivo",
        ),
        date_columns=("Fecha",),
    )
    product_rows = [
        (
            sheet.route_date,
            sheet.id,
            sheet.seller_name,
            item.item_label,
            item.item_category,
            item.total_out,
            item.total_returned,
            item.physical_sold,
            item.registered_sold,
            item.physical_difference,
            STATUS_LABELS[sheet.status],
        )
        for sheet in sheets
        for item in settlements[sheet.id].products
    ]
    add_sheet(
        workbook,
        "Diferencias de producto",
        (
            "Fecha",
            "Ruta ID",
            "Vendedor",
            "Producto",
            "Categoría",
            "Salido",
            "Regresado",
            "Vendido físico",
            "Vendido registrado",
            "Diferencia",
            "Estado ruta",
        ),
        _rows_or_message(product_rows, 11, "Sin movimientos de producto"),
        quantity_columns=(
            "Salido",
            "Regresado",
            "Vendido físico",
            "Vendido registrado",
            "Diferencia",
        ),
        date_columns=("Fecha",),
    )
    return workbook


def _list_filtered_sheets(
    db: Session,
    *,
    start_date: date | None,
    end_date: date | None,
    seller: str | None = None,
    route_label: str | None = None,
    status: str | None = None,
    period_field: str,
) -> list[SimpleRouteSheet]:
    _validate_period(start_date, end_date)
    _validate_status(status)
    sheets = simple_route_repository.list_sheets(
        db,
        date_from=start_date if period_field == "route" else None,
        date_to=end_date if period_field == "route" else None,
        status=status,
    )
    normalized_seller = _normalize_text_filter(seller)
    normalized_route = _normalize_text_filter(route_label)
    filtered = [
        sheet
        for sheet in sheets
        if _matches_text(sheet.seller_name, normalized_seller)
        and _matches_text(sheet.route_label, normalized_route)
    ]
    if period_field == "events":
        filtered = [sheet for sheet in filtered if sheet.status != "cancelled"]
    return filtered


def _add_rankings_sheet(
    workbook: Workbook,
    sheets: list[SimpleRouteSheet],
    settlements: dict[int, Any],
) -> None:
    totals: dict[tuple[str, str], dict[str, Any]] = defaultdict(
        lambda: {
            "routes": 0,
            "sales": ZERO_MONEY,
            "delivered": ZERO_MONEY,
            "difference": ZERO_MONEY,
        }
    )
    for sheet in sheets:
        settlement = settlements[sheet.id]
        for group, label in (
            ("Vendedor", sheet.seller_name),
            ("Ruta", sheet.route_label or "Sin ruta / zona"),
        ):
            row = totals[(group, label)]
            row["routes"] += 1
            row["sales"] += settlement.route_sales_total
            row["delivered"] += settlement.cash_delivered or ZERO_MONEY
            row["difference"] += settlement.cash_difference or ZERO_MONEY
    rows = [
        (
            group,
            label,
            values["routes"],
            values["sales"],
            values["delivered"],
            values["difference"],
        )
        for (group, label), values in sorted(
            totals.items(),
            key=lambda item: (item[0][0], -item[1]["sales"], item[0][1].casefold()),
        )
    ]
    add_sheet(
        workbook,
        "Top vendedores y rutas",
        (
            "Tipo",
            "Vendedor / ruta",
            "Cantidad de rutas",
            "Ventas",
            "Efectivo entregado",
            "Diferencia",
        ),
        _rows_or_message(rows, 6, EMPTY_PERIOD_MESSAGE),
        money_columns=("Ventas", "Efectivo entregado", "Diferencia"),
    )


def _sale_rows(sheet: SimpleRouteSheet) -> list[tuple[Any, ...]]:
    rows: list[tuple[Any, ...]] = []
    for sale in sheet.sales:
        collected = sale.total_amount - sale.remaining_amount
        for item in sale.items:
            rows.append(
                (
                    sale.sale_date,
                    sale.id,
                    sale.customer_name_snapshot,
                    SALE_TYPE_LABELS[sale.sale_type],
                    item.product_name_snapshot,
                    item.item_category,
                    item.quantity,
                    item.unit_price,
                    item.subtotal,
                    sale.total_amount,
                    collected,
                    sale.remaining_amount,
                    "Activa" if sale.status == "active" else "Anulada",
                    sale.notes,
                )
            )
    return rows


def _rows_or_message(
    rows: list[tuple[Any, ...]],
    column_count: int,
    message: str,
) -> list[tuple[Any, ...]]:
    if rows:
        return rows
    return [(message, *([None] * (column_count - 1)))]


def _date_in_period(
    value: date,
    start_date: date | None,
    end_date: date | None,
) -> bool:
    return (start_date is None or value >= start_date) and (
        end_date is None or value <= end_date
    )


def _matches_text(value: str | None, normalized_filter: str | None) -> bool:
    if normalized_filter is None:
        return True
    return normalized_filter in (value or "").strip().casefold()


def _normalize_text_filter(value: str | None) -> str | None:
    normalized = (value or "").strip().casefold()
    return normalized or None


def _validate_period(start_date: date | None, end_date: date | None) -> None:
    if start_date is not None and end_date is not None and start_date > end_date:
        raise ValueError("La fecha inicial no puede ser posterior a la fecha final.")


def _validate_status(status: str | None) -> None:
    if status is not None and status not in STATUS_LABELS:
        raise ValueError("El estado de Ruta Simple no es válido.")


def _validate_sale_type(sale_type: str | None) -> None:
    if sale_type is not None and sale_type not in {*SALE_TYPE_LABELS, "all"}:
        raise ValueError("El tipo de venta de Ruta Simple no es válido.")


def _sum_money(values) -> Decimal:
    return sum(values, start=ZERO_MONEY).quantize(Decimal("0.01"))
