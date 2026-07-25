from collections import defaultdict
from collections.abc import Iterable, Sequence
from datetime import date, datetime
from decimal import Decimal
from enum import Enum
from io import BytesIO
from typing import Any

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.services import credit_payments as credit_payment_service
from app.services import inventory as inventory_service
from app.services import operating_expenses as operating_expense_service
from app.services import payroll as payroll_service
from app.services import production as production_service
from app.services import production_material_consumptions as consumption_service
from app.services import production_recipes as recipe_service
from app.services import purchases as purchase_service
from app.services import receivables as receivable_service
from app.services import route_cash_receipts as receipt_service
from app.services import route_closures as closure_service
from app.services import route_loads as route_load_service
from app.services import route_returns as return_service
from app.services import sales as sale_service
from app.services.management_dashboard import get_management_dashboard
from app.services.monthly_operational_summary import (
    get_monthly_operational_summary,
)

EXCEL_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
MONEY_FORMAT = '"$"#,##0.00'
QUANTITY_FORMAT = "#,##0.000"
DATE_FORMAT = "yyyy-mm-dd"
DATETIME_FORMAT = "yyyy-mm-dd hh:mm"
LOW_STOCK_THRESHOLD = Decimal("5.000")
RECENT_MOVEMENTS_LIMIT = 200
HEADER_FILL = PatternFill(fill_type="solid", fgColor="0F766E")
HEADER_FONT = Font(bold=True, color="FFFFFF")

LABELS = {
    "active_locations_count": "Ubicaciones activas",
    "active_products_with_stock_count": "Productos activos con stock",
    "active_purchases_count": "Compras activas",
    "active_sales_count": "Ventas activas",
    "alert_threshold": "Umbral de alerta",
    "average_daily_counted": "Promedio diario contado",
    "average_sale_amount": "Venta promedio",
    "cash_final_total": "Caja final",
    "cash_sales_count": "Ventas de contado",
    "cash_sales_total": "Ventas de contado",
    "cash_sessions_count": "Cuadres de caja",
    "cash_total_counted": "Total contado",
    "cash_total_expenses_after_count": "Gastos posteriores al conteo",
    "closed_cash_sessions_count": "Cuadres cerrados",
    "closures_balanced_count": "Cierres cuadrados",
    "closures_with_shortage_count": "Cierres con faltante",
    "closures_with_surplus_count": "Cierres con sobrante",
    "credit_sales_count": "Ventas a crédito",
    "credit_sales_created_total": "Crédito generado",
    "credit_sales_total": "Ventas a crédito",
    "customers_with_balance_count": "Clientes con saldo",
    "customers_with_credit_sales_count": "Clientes con ventas a crédito",
    "days_in_month": "Días del mes",
    "days_with_alerts_count": "Días con alertas",
    "days_with_cash_session": "Días con cuadre",
    "days_without_cash_session": "Días sin cuadre",
    "deductions_total": "Deducciones",
    "delivered_cash_total": "Efectivo entregado",
    "difference_total": "Diferencia",
    "gross_pay_total": "Nómina bruta",
    "material_consumptions_count": "Consumos de insumos",
    "material_consumptions_total_cost": "Costo de insumos",
    "net_pay_total": "Nómina neta",
    "open_cash_sessions_count": "Cuadres abiertos",
    "operating_expenses_total": "Gastos operativos",
    "overdue_balance_total": "Cartera vencida",
    "overdue_customers_count": "Clientes vencidos",
    "paid_expenses_total": "Gastos pagados",
    "paid_purchases_total": "Compras pagadas",
    "payments_count": "Cobros",
    "payments_total": "Cobros recibidos",
    "payroll_runs_count": "Roles de pago",
    "pending_balance_total": "Cartera pendiente",
    "pending_credit_total": "Crédito pendiente",
    "production_applied_to_inventory_count": "Aplicaciones a inventario",
    "production_batches_count": "Producciones",
    "production_items_total_quantity": "Cantidad producida",
    "production_waste_total_quantity": "Merma",
    "purchases_total": "Compras",
    "received_amount_total": "Efectivo recibido",
    "receipts_count": "Recibos",
    "route_cash_receipts_total": "Efectivo de rutas en caja",
    "route_closures_count": "Cierres de ruta",
    "route_delivered_cash_total": "Efectivo entregado por rutas",
    "route_difference_total": "Diferencia de rutas",
    "route_expected_cash_total": "Efectivo esperado de rutas",
    "route_loads_count": "Cargas de ruta",
    "route_returns_count": "Retornos de ruta",
    "sales_total": "Ventas totales",
    "total_alerts": "Alertas",
    "total_bills": "Billetes",
    "total_business_cash_inflow_reference": "Ingreso de efectivo de referencia",
    "total_coins": "Monedas",
    "total_counted": "Total contado",
    "total_expenses": "Gastos de caja",
    "total_final": "Total final",
    "total_quantity_loaded": "Cantidad cargada",
    "total_quantity_net": "Producción neta",
    "total_quantity_produced": "Cantidad producida",
    "total_quantity_returned": "Cantidad retornada",
    "total_quantity_waste": "Merma",
    "total_sales_amount": "Ventas totales",
    "total_stock_quantity": "Stock total",
    "unpaid_expenses_total": "Gastos pendientes",
    "unpaid_purchases_total": "Compras pendientes",
    "voided_expenses_count": "Gastos anulados",
    "voided_payments_count": "Cobros anulados",
    "voided_payroll_runs_count": "Roles anulados",
    "voided_purchases_count": "Compras anuladas",
    "voided_sales_count": "Ventas anuladas",
}


def create_workbook() -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    return workbook


def safe_text(value: Any) -> str:
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value or ""))
    if text.startswith(("=", "+", "-", "@", "\t", "\r")):
        return f"'{text}"
    return text


def decimal_to_number(value: Decimal | None) -> Decimal | None:
    return value


def add_sheet(
    workbook: Workbook,
    title: str,
    headers: Sequence[str],
    rows: Iterable[Sequence[Any]],
    *,
    money_columns: Sequence[str] = (),
    quantity_columns: Sequence[str] = (),
    date_columns: Sequence[str] = (),
    datetime_columns: Sequence[str] = (),
):
    sheet = workbook.create_sheet(_safe_sheet_title(title))
    normalized_headers = [safe_text(header) for header in headers]
    sheet.append(normalized_headers)

    row_count = 0
    for row in rows:
        values = [_excel_value(value) for value in row]
        if len(values) < len(normalized_headers):
            values.extend([None] * (len(normalized_headers) - len(values)))
        sheet.append(values[: len(normalized_headers)])
        row_count += 1

    if row_count == 0:
        sheet.append(["Sin datos", *([None] * (len(normalized_headers) - 1))])

    format_header(sheet)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(normalized_headers))}{sheet.max_row}"
    )
    _format_columns(sheet, normalized_headers, money_columns, MONEY_FORMAT)
    _format_columns(
        sheet,
        normalized_headers,
        quantity_columns,
        QUANTITY_FORMAT,
    )
    _format_columns(sheet, normalized_headers, date_columns, DATE_FORMAT)
    _format_columns(
        sheet,
        normalized_headers,
        datetime_columns,
        DATETIME_FORMAT,
    )
    autosize_columns(sheet)
    return sheet


def autosize_columns(sheet) -> None:
    for column_cells in sheet.iter_cols(
        min_row=1,
        max_row=sheet.max_row,
        max_col=sheet.max_column,
    ):
        width = max(
            (_display_length(cell.value) for cell in column_cells),
            default=0,
        )
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = (
            min(max(width + 2, 12), 48)
        )


def format_header(sheet) -> None:
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 24


def build_excel_response(
    workbook: Workbook,
    filename: str,
) -> StreamingResponse:
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type=EXCEL_MEDIA_TYPE,
        headers={
            "Content-Disposition": (
                f'attachment; filename="{_safe_filename(filename)}"'
            ),
            "Cache-Control": "no-store",
        },
    )


def build_management_dashboard_workbook(
    db: Session,
    date_from: date | None = None,
    date_to: date | None = None,
) -> Workbook:
    report = get_management_dashboard(db, date_from=date_from, date_to=date_to)
    workbook = create_workbook()

    summary_rows: list[tuple[Any, ...]] = [
        ("Período", "Fecha desde", report.period.date_from),
        ("Período", "Fecha hasta", report.period.date_to),
    ]
    for section, values in (
        ("Caja", report.cash.model_dump()),
        ("Ventas", report.sales.model_dump()),
        ("Cartera", report.receivables.model_dump()),
        ("Gastos operativos", report.operating_expenses.model_dump()),
        ("Compras", report.purchases.model_dump()),
        ("Producción", report.production.model_dump()),
        ("Inventario", report.inventory.model_dump()),
        ("Rutas", report.routes.model_dump()),
        ("Nómina", report.payroll.model_dump()),
    ):
        summary_rows.extend(_metric_rows(section, values))
    add_sheet(
        workbook,
        "Resumen",
        ("Sección", "Indicador", "Valor"),
        summary_rows,
    )
    _add_metric_sheet(workbook, "Caja", report.cash.model_dump())
    _add_metric_sheet(workbook, "Ventas", report.sales.model_dump())
    _add_metric_sheet(workbook, "Cartera", report.receivables.model_dump())

    expense_rows = list(
        _metric_rows(
            "Resumen",
            report.operating_expenses.model_dump(),
        )
    )
    expense_rows.extend(
        (
            "Categoría",
            item.category_name,
            item.total_amount,
        )
        for item in report.operating_expenses.top_categories
    )
    add_sheet(
        workbook,
        "Gastos operativos",
        ("Sección", "Indicador", "Valor"),
        expense_rows,
    )

    purchase_rows = list(_metric_rows("Resumen", report.purchases.model_dump()))
    purchase_rows.extend(
        ("Proveedor", item.supplier_name, item.total_amount)
        for item in report.purchases.top_suppliers
    )
    add_sheet(
        workbook,
        "Compras",
        ("Sección", "Indicador", "Valor"),
        purchase_rows,
    )
    _add_metric_sheet(workbook, "Producción", report.production.model_dump())

    inventory_rows = list(
        _metric_rows("Resumen", report.inventory.model_dump())
    )
    inventory_rows.extend(
        (
            "Stock bajo",
            item.product_name,
            item.total_stock_quantity,
        )
        for item in report.inventory.low_stock_items
    )
    add_sheet(
        workbook,
        "Inventario",
        ("Sección", "Indicador", "Valor"),
        inventory_rows,
    )
    _add_metric_sheet(workbook, "Rutas", report.routes.model_dump())
    _add_metric_sheet(workbook, "Nómina", report.payroll.model_dump())
    add_sheet(
        workbook,
        "Alertas",
        ("Tipo", "Título", "Mensaje", "Cantidad"),
        (
            (alert.type, alert.title, alert.message, alert.count)
            for alert in report.alerts
        ),
    )

    series: dict[date, dict[str, Any]] = defaultdict(dict)
    for point in report.charts.daily_sales:
        series[point.date].update(point.model_dump(exclude={"date"}))
    for point in report.charts.daily_cash:
        series[point.date].update(point.model_dump(exclude={"date"}))
    for point in report.charts.daily_expenses:
        series[point.date].update(point.model_dump(exclude={"date"}))
    add_sheet(
        workbook,
        "Series diarias",
        (
            "Fecha",
            "Ventas contado",
            "Ventas crédito",
            "Cobros",
            "Caja contada",
            "Caja final",
            "Gastos operativos",
            "Gastos de caja",
        ),
        (
            (
                report_date,
                values.get("cash_sales_total"),
                values.get("credit_sales_total"),
                values.get("payments_total"),
                values.get("cash_total_counted"),
                values.get("cash_final_total"),
                values.get("operating_expenses_total"),
                values.get("cash_expenses_total"),
            )
            for report_date, values in sorted(series.items())
        ),
        money_columns=(
            "Ventas contado",
            "Ventas crédito",
            "Cobros",
            "Caja contada",
            "Caja final",
            "Gastos operativos",
            "Gastos de caja",
        ),
        date_columns=("Fecha",),
    )
    add_sheet(
        workbook,
        "Producción producto",
        ("Producto", "Cantidad producida"),
        (
            (point.product_name, point.quantity_produced)
            for point in report.charts.production_by_product
        ),
        quantity_columns=("Cantidad producida",),
    )
    return workbook


def build_monthly_summary_workbook(
    db: Session,
    year: int,
    month: int,
) -> Workbook:
    report = get_monthly_operational_summary(db, year=year, month=month)
    workbook = create_workbook()
    add_sheet(
        workbook,
        "Resumen mensual",
        ("Sección", "Indicador", "Valor"),
        [
            ("Período", "Año", report.period.year),
            ("Período", "Mes", report.period.month),
            ("Período", "Fecha desde", report.period.date_from),
            ("Período", "Fecha hasta", report.period.date_to),
            *list(_metric_rows("Totales", report.totals.model_dump())),
        ],
    )
    _add_metric_sheet(workbook, "Caja", report.cash_summary.model_dump())

    sales_rows = list(_metric_rows("Resumen", report.sales_summary.model_dump()))
    sales_rows.extend(
        ("Producto vendido", item.product_name_snapshot, item.amount)
        for item in report.sales_summary.products_sold
    )
    add_sheet(
        workbook,
        "Ventas",
        ("Sección", "Indicador", "Valor"),
        sales_rows,
    )
    _add_metric_sheet(
        workbook,
        "Cartera",
        report.receivables_summary.model_dump(),
    )

    production_rows = list(
        _metric_rows("Resumen", report.production_summary.model_dump())
    )
    production_rows.extend(
        ("Producto", item.product_name_snapshot, item.quantity_net)
        for item in report.production_summary.products
    )
    add_sheet(
        workbook,
        "Producción",
        ("Sección", "Indicador", "Valor"),
        production_rows,
    )

    route_rows = [
        *list(
            _metric_rows(
                "Operación",
                report.route_operations_summary.model_dump(),
            )
        ),
        *list(
            _metric_rows(
                "Finanzas",
                report.route_financial_summary.model_dump(),
            )
        ),
    ]
    add_sheet(
        workbook,
        "Rutas",
        ("Sección", "Indicador", "Valor"),
        route_rows,
    )
    _add_metric_sheet(
        workbook,
        "Recibos",
        report.route_cash_receipts_summary.model_dump(),
    )
    add_sheet(
        workbook,
        "Alertas",
        ("Código", "Mensaje", "Cantidad"),
        (
            (item.code, item.message, item.count)
            for item in report.alerts_summary.alerts_by_code
        ),
    )
    add_sheet(
        workbook,
        "Días del mes",
        (
            "Fecha",
            "Existe caja",
            "Total contado",
            "Gastos de caja",
            "Total final",
            "Producción neta",
            "Ventas contado",
            "Ventas crédito",
            "Cobros",
            "Efectivo rutas",
            "Efectivo registrado",
            "Alertas",
            "Códigos de alerta",
        ),
        (
            (
                row.date,
                row.cash_session_exists,
                row.cash_total_counted,
                row.cash_total_expenses,
                row.cash_total_final,
                row.production_quantity_net,
                row.cash_sales_total,
                row.credit_sales_total,
                row.payments_total,
                row.route_delivered_cash_total,
                row.route_cash_registered_total,
                row.alerts_count,
                ", ".join(row.alert_codes),
            )
            for row in report.daily_rows
        ),
        money_columns=(
            "Total contado",
            "Gastos de caja",
            "Total final",
            "Ventas contado",
            "Ventas crédito",
            "Cobros",
            "Efectivo rutas",
            "Efectivo registrado",
        ),
        quantity_columns=("Producción neta",),
        date_columns=("Fecha",),
    )
    return workbook


def build_sales_workbook(
    db: Session,
    *,
    date_from: date | None = None,
    date_to: date | None = None,
    sale_type: str | None = None,
    status: str | None = None,
    customer_id: int | None = None,
    route_load_id: int | None = None,
) -> Workbook:
    sales = sale_service.list_sales(
        db=db,
        date_from=date_from,
        date_to=date_to,
        sale_type=sale_type,
        status=status,
        customer_id=customer_id,
        route_load_id=route_load_id,
    )
    active_sales = [sale for sale in sales if sale.status == "active"]
    workbook = create_workbook()
    add_sheet(
        workbook,
        "Resumen",
        ("Indicador", "Valor"),
        (
            ("Ventas listadas", len(sales)),
            ("Ventas activas", len(active_sales)),
            (
                "Total activo",
                sum(
                    (sale.total_amount for sale in active_sales),
                    start=Decimal("0.00"),
                ),
            ),
            (
                "Saldo activo",
                sum(
                    (sale.balance_amount for sale in active_sales),
                    start=Decimal("0.00"),
                ),
            ),
        ),
    )
    add_sheet(
        workbook,
        "Ventas",
        (
            "Fecha",
            "Número de venta",
            "Cliente",
            "Tipo de venta",
            "Método de pago",
            "Estado",
            "Estado crédito",
            "Total",
            "Pagado",
            "Saldo",
            "Vencimiento",
            "Carga",
            "Ruta",
            "Notas",
            "Motivo anulación",
        ),
        (
            (
                sale.sale_date,
                sale.sale_number,
                sale.customer_name_snapshot,
                sale.sale_type,
                sale.payment_method,
                sale.status,
                sale.credit_status,
                sale.total_amount,
                sale.paid_amount,
                sale.balance_amount,
                sale.due_date,
                sale.route_load_number,
                sale.route_name,
                sale.notes,
                sale.void_reason,
            )
            for sale in sales
        ),
        money_columns=("Total", "Pagado", "Saldo"),
        date_columns=("Fecha", "Vencimiento"),
    )
    add_sheet(
        workbook,
        "Items",
        (
            "Venta",
            "SKU",
            "Producto",
            "Presentación",
            "Cantidad",
            "Precio unitario",
            "Subtotal",
            "Origen de precio",
            "Lista de precios",
            "Nota de precio",
        ),
        (
            (
                sale.sale_number,
                item.product_sku_snapshot,
                item.product_name_snapshot,
                item.product_presentation_snapshot,
                item.quantity,
                item.unit_price,
                item.subtotal,
                item.price_source,
                item.price_list_name_snapshot,
                item.pricing_note,
            )
            for sale in sales
            for item in sale.items
        ),
        money_columns=("Precio unitario", "Subtotal"),
        quantity_columns=("Cantidad",),
    )
    return workbook


def build_receivables_workbook(
    db: Session,
    *,
    date_from: date | None = None,
    date_to: date | None = None,
    customer_id: int | None = None,
) -> Workbook:
    customers = receivable_service.list_receivable_customers(
        db=db,
        only_with_balance=True,
        date_from=date_from,
        date_to=date_to,
    )
    if customer_id is not None:
        customers = [
            customer
            for customer in customers
            if customer.customer_id == customer_id
        ]

    if customer_id is not None:
        customer_statement = receivable_service.get_customer_statement(
            db=db,
            customer_id=customer_id,
            date_from=date_from,
            date_to=date_to,
            include_paid=True,
        )
        statements = [customer_statement]
        summary_values = (
            int(customer_statement.total_balance > Decimal("0.00")),
            customer_statement.total_credit_sales,
            customer_statement.total_paid,
            customer_statement.total_balance,
        )
    else:
        summary = receivable_service.get_receivables_summary(
            db,
            date_from=date_from,
            date_to=date_to,
        )
        statements = [
            receivable_service.get_customer_statement(
                db=db,
                customer_id=customer.customer_id,
                date_from=date_from,
                date_to=date_to,
                include_paid=True,
            )
            for customer in customers
        ]
        summary_values = (
            summary.total_customers_with_balance,
            summary.total_credit_sales,
            summary.total_paid,
            summary.total_balance,
        )
    overdue = receivable_service.list_overdue_receivables(db)
    overdue = [
        item
        for item in overdue
        if (customer_id is None or item.customer_id == customer_id)
        and (date_from is None or item.sale_date >= date_from)
        and (date_to is None or item.sale_date <= date_to)
    ]

    workbook = create_workbook()
    add_sheet(
        workbook,
        "Resumen",
        ("Indicador", "Valor"),
        (
            ("Clientes con saldo", summary_values[0]),
            ("Ventas a crédito", summary_values[1]),
            ("Total pagado", summary_values[2]),
            ("Saldo pendiente", summary_values[3]),
        ),
    )
    add_sheet(
        workbook,
        "Clientes con saldo",
        (
            "Cliente",
            "Teléfono",
            "Sector",
            "Activo",
            "Total vendido a crédito",
            "Total pagado",
            "Saldo pendiente",
            "Ventas pendientes",
        ),
        (
            (
                customer.customer_name,
                customer.primary_phone,
                customer.sector,
                customer.is_active,
                customer.total_credit_sales,
                customer.total_paid,
                customer.total_balance,
                customer.pending_sales_count,
            )
            for customer in customers
        ),
        money_columns=(
            "Total vendido a crédito",
            "Total pagado",
            "Saldo pendiente",
        ),
    )
    statement_rows: list[tuple[Any, ...]] = []
    for statement in statements:
        statement_rows.extend(
            (
                statement.customer_name,
                "Venta",
                sale.sale_date,
                sale.sale_number,
                sale.total_amount,
                sale.paid_amount,
                sale.balance_amount,
                sale.credit_status,
                sale.status,
            )
            for sale in statement.credit_sales
        )
        statement_rows.extend(
            (
                statement.customer_name,
                "Abono",
                payment.payment_date,
                payment.sale_number,
                payment.amount,
                None,
                None,
                payment.payment_method,
                payment.status,
            )
            for payment in statement.payments
        )
    add_sheet(
        workbook,
        "Estado de cuenta",
        (
            "Cliente",
            "Tipo",
            "Fecha",
            "Venta",
            "Importe",
            "Pagado",
            "Saldo",
            "Detalle",
            "Estado",
        ),
        statement_rows,
        money_columns=("Importe", "Pagado", "Saldo"),
        date_columns=("Fecha",),
    )
    add_sheet(
        workbook,
        "Vencidos",
        (
            "Cliente",
            "Teléfono",
            "Sector",
            "Venta",
            "Fecha venta",
            "Vencimiento",
            "Total",
            "Pagado",
            "Saldo",
        ),
        (
            (
                item.customer_name,
                item.primary_phone,
                item.sector,
                item.sale_number,
                item.sale_date,
                item.due_date,
                item.total_amount,
                item.paid_amount,
                item.balance_amount,
            )
            for item in overdue
        ),
        money_columns=("Total", "Pagado", "Saldo"),
        date_columns=("Fecha venta", "Vencimiento"),
    )
    return workbook


def build_operating_expenses_workbook(
    db: Session,
    *,
    date_from: date | None = None,
    date_to: date | None = None,
    category_id: int | None = None,
    supplier_id: int | None = None,
    payment_status: str | None = None,
    status: str | None = None,
) -> Workbook:
    expenses = operating_expense_service.list_expenses(
        db=db,
        date_from=date_from,
        date_to=date_to,
        category_id=category_id,
        supplier_id=supplier_id,
        payment_status=payment_status,
        status=status,
    )
    active_expenses = [
        expense for expense in expenses if expense.status == "active"
    ]
    workbook = create_workbook()
    add_sheet(
        workbook,
        "Resumen",
        ("Indicador", "Valor"),
        (
            ("Gastos listados", len(expenses)),
            ("Gastos activos", len(active_expenses)),
            (
                "Total activo",
                sum(
                    (expense.amount for expense in active_expenses),
                    start=Decimal("0.00"),
                ),
            ),
        ),
    )
    add_sheet(
        workbook,
        "Gastos operativos",
        (
            "Fecha",
            "Número",
            "Categoría",
            "Proveedor",
            "Descripción",
            "Monto",
            "Estado de pago",
            "Método de pago",
            "Estado",
            "Documento",
            "Fecha vencimiento",
            "Fecha pago",
            "Notas",
            "Motivo anulación",
        ),
        (
            (
                expense.expense_date,
                expense.expense_number,
                expense.category.name,
                expense.supplier.name if expense.supplier else None,
                expense.description,
                expense.amount,
                expense.payment_status,
                expense.payment_method,
                expense.status,
                expense.document_number,
                expense.due_date,
                expense.paid_date,
                expense.notes,
                expense.void_reason,
            )
            for expense in expenses
        ),
        money_columns=("Monto",),
        date_columns=(
            "Fecha",
            "Fecha vencimiento",
            "Fecha pago",
        ),
    )
    return workbook


def build_purchases_workbook(
    db: Session,
    *,
    date_from: date | None = None,
    date_to: date | None = None,
    supplier_id: int | None = None,
    payment_status: str | None = None,
    status: str | None = None,
) -> Workbook:
    purchases = purchase_service.list_purchases(
        db=db,
        date_from=date_from,
        date_to=date_to,
        supplier_id=supplier_id,
        payment_status=payment_status,
        status=status,
    )
    workbook = create_workbook()
    add_sheet(
        workbook,
        "Compras",
        (
            "Fecha",
            "Número compra",
            "Proveedor",
            "Documento",
            "Tipo documento",
            "Estado",
            "Estado de pago",
            "Subtotal",
            "Descuento",
            "Impuesto",
            "Total",
            "Notas",
            "Motivo anulación",
        ),
        (
            (
                purchase.purchase_date,
                purchase.purchase_number,
                purchase.supplier.name,
                purchase.document_number,
                purchase.document_type,
                purchase.status,
                purchase.payment_status,
                purchase.subtotal_amount,
                purchase.discount_amount,
                purchase.tax_amount,
                purchase.total_amount,
                purchase.notes,
                purchase.void_reason,
            )
            for purchase in purchases
        ),
        money_columns=(
            "Subtotal",
            "Descuento",
            "Impuesto",
            "Total",
        ),
        date_columns=("Fecha",),
    )
    add_sheet(
        workbook,
        "Items",
        (
            "Compra",
            "SKU",
            "Producto",
            "Descripción",
            "Cantidad",
            "Costo unitario",
            "Descuento",
            "Impuesto",
            "Subtotal",
            "Total",
            "Notas",
        ),
        (
            (
                purchase.purchase_number,
                item.product_sku_snapshot,
                item.product_name_snapshot,
                item.description,
                item.quantity,
                item.unit_cost,
                item.discount_amount,
                item.tax_amount,
                item.subtotal_amount,
                item.total_amount,
                item.notes,
            )
            for purchase in purchases
            for item in purchase.items
        ),
        money_columns=(
            "Costo unitario",
            "Descuento",
            "Impuesto",
            "Subtotal",
            "Total",
        ),
        quantity_columns=("Cantidad",),
    )
    return workbook


def build_inventory_workbook(
    db: Session,
    *,
    location_id: int | None = None,
    product_id: int | None = None,
    low_stock_only: bool = False,
) -> Workbook:
    stocks = inventory_service.list_stocks(
        db=db,
        location_id=location_id,
        product_id=product_id,
    )
    if low_stock_only:
        stocks = [
            stock
            for stock in stocks
            if stock.quantity_on_hand <= LOW_STOCK_THRESHOLD
        ]
    locations = inventory_service.list_locations(db)
    if location_id is not None:
        locations = [
            location for location in locations if location.id == location_id
        ]
    movements = inventory_service.list_movements(
        db=db,
        location_id=location_id,
        product_id=product_id,
    )[:RECENT_MOVEMENTS_LIMIT]

    workbook = create_workbook()
    add_sheet(
        workbook,
        "Stock actual",
        (
            "Producto",
            "SKU",
            "Tipo productivo",
            "Ubicación",
            "Cantidad",
            "Última actualización",
        ),
        (
            (
                stock.product.name,
                stock.product.sku,
                stock.product.product_kind,
                stock.location.name,
                stock.quantity_on_hand,
                stock.updated_at,
            )
            for stock in stocks
        ),
        quantity_columns=("Cantidad",),
        datetime_columns=("Última actualización",),
    )
    add_sheet(
        workbook,
        "Ubicaciones",
        (
            "Código",
            "Nombre",
            "Tipo",
            "Activa",
            "Descripción",
            "Notas",
        ),
        (
            (
                location.code,
                location.name,
                location.location_type,
                location.is_active,
                location.description,
                location.notes,
            )
            for location in locations
        ),
    )
    add_sheet(
        workbook,
        "Movimientos recientes",
        (
            "Fecha",
            "Tipo",
            "Producto",
            "SKU",
            "Ubicación",
            "Cantidad",
            "Costo unitario",
            "Costo total",
            "Estado",
            "Motivo",
            "Motivo anulación",
        ),
        (
            (
                movement.movement_date,
                movement.movement_type,
                movement.product_name_snapshot,
                movement.product_sku_snapshot,
                movement.location.name,
                movement.quantity,
                movement.unit_cost,
                movement.total_cost,
                movement.status,
                movement.reason,
                movement.void_reason,
            )
            for movement in movements
        ),
        money_columns=("Costo unitario", "Costo total"),
        quantity_columns=("Cantidad",),
        date_columns=("Fecha",),
    )
    return workbook


def build_production_workbook(
    db: Session,
    *,
    date_from: date | None = None,
    date_to: date | None = None,
    status: str | None = None,
    product_id: int | None = None,
) -> Workbook:
    batches = production_service.list_production_batches(
        db=db,
        date_from=date_from,
        date_to=date_to,
        status=status,
        product_id=product_id,
    )
    batch_ids = {batch.id for batch in batches}
    consumptions = consumption_service.list_material_consumptions(
        db=db,
        date_from=date_from,
        date_to=date_to,
        status=status,
    )
    consumptions = [
        item
        for item in consumptions
        if item.production_batch_id in batch_ids
    ]
    recipes = recipe_service.list_recipes(
        db=db,
        output_product_id=product_id,
    )
    workbook = create_workbook()
    add_sheet(
        workbook,
        "Producciones",
        ("Fecha", "Número lote", "Estado", "Notas", "Motivo anulación"),
        (
            (
                batch.production_date,
                batch.production_number,
                batch.status,
                batch.notes,
                batch.void_reason,
            )
            for batch in batches
        ),
        date_columns=("Fecha",),
    )
    add_sheet(
        workbook,
        "Items producidos",
        (
            "Producción",
            "SKU",
            "Producto",
            "Presentación",
            "Cantidad producida",
            "Merma",
            "Cantidad neta",
            "Notas",
        ),
        (
            (
                batch.production_number,
                item.product_sku_snapshot,
                item.product_name_snapshot,
                item.product_presentation_snapshot,
                item.quantity_produced,
                item.quantity_waste,
                item.quantity_produced - item.quantity_waste,
                item.notes,
            )
            for batch in batches
            for item in batch.items
        ),
        quantity_columns=(
            "Cantidad producida",
            "Merma",
            "Cantidad neta",
        ),
    )
    add_sheet(
        workbook,
        "Aplicaciones inventario",
        (
            "Producción",
            "Ubicación",
            "Estado",
            "Fecha aplicación",
            "Producto",
            "Cantidad aplicada",
            "Motivo anulación",
        ),
        (
            (
                batch.production_number,
                application.location.name,
                application.status,
                application.applied_at,
                item.product.name,
                item.quantity_applied,
                application.void_reason,
            )
            for batch in batches
            for application in batch.inventory_applications
            for item in application.items
        ),
        quantity_columns=("Cantidad aplicada",),
        datetime_columns=("Fecha aplicación",),
    )
    add_sheet(
        workbook,
        "Consumo de insumos",
        (
            "Fecha",
            "Consumo",
            "Producción",
            "Ubicación origen",
            "Estado",
            "Insumo",
            "SKU",
            "Cantidad base",
            "Cantidad consumida",
            "Costo unitario",
            "Costo total",
            "Fórmula",
            "Notas",
        ),
        (
            (
                consumption.movement_date,
                consumption.consumption_number,
                consumption.production_batch.production_number,
                consumption.source_location.name,
                consumption.status,
                item.input_product_name_snapshot,
                item.input_product_sku_snapshot,
                item.quantity_required_base,
                item.quantity_consumed,
                item.unit_cost,
                item.total_cost,
                item.recipe_name_snapshot,
                consumption.notes,
            )
            for consumption in consumptions
            for item in consumption.items
        ),
        money_columns=("Costo unitario", "Costo total"),
        quantity_columns=("Cantidad base", "Cantidad consumida"),
        date_columns=("Fecha",),
    )
    add_sheet(
        workbook,
        "Fórmulas",
        (
            "Número",
            "Nombre",
            "Producto salida",
            "Cantidad salida",
            "Estado",
            "Vigente desde",
            "Vigente hasta",
            "Insumo",
            "Cantidad requerida",
            "Merma %",
            "Costo total",
            "Activo",
        ),
        (
            (
                recipe.recipe_number,
                recipe.name,
                recipe.output_product.name,
                recipe.output_quantity,
                recipe.status,
                recipe.valid_from,
                recipe.valid_to,
                item.input_product_name_snapshot,
                item.quantity_required,
                item.waste_percentage,
                item.total_cost,
                item.is_active,
            )
            for recipe in recipes
            for item in recipe.items
        ),
        money_columns=("Costo total",),
        quantity_columns=("Cantidad salida", "Cantidad requerida"),
        date_columns=("Vigente desde", "Vigente hasta"),
    )
    return workbook


def build_routes_workbook(
    db: Session,
    *,
    date_from: date | None = None,
    date_to: date | None = None,
    route_load_id: int | None = None,
    status: str | None = None,
) -> Workbook:
    loads = route_load_service.list_route_loads(
        db=db,
        date_from=date_from,
        date_to=date_to,
        status=status,
    )
    if route_load_id is not None:
        loads = [load for load in loads if load.id == route_load_id]
    returns = return_service.list_route_returns(
        db=db,
        date_from=date_from,
        date_to=date_to,
        status=status,
        route_load_id=route_load_id,
    )
    closures = closure_service.list_route_closures(
        db=db,
        date_from=date_from,
        date_to=date_to,
        status=status,
        route_load_id=route_load_id,
    )
    receipts = receipt_service.list_route_cash_receipts(
        db=db,
        date_from=date_from,
        date_to=date_to,
        status=status,
    )
    if route_load_id is not None:
        receipts = [
            receipt
            for receipt in receipts
            if receipt.route_load_id == route_load_id
        ]
    sales = sale_service.list_sales(
        db=db,
        date_from=date_from,
        date_to=date_to,
        route_load_id=route_load_id,
    )
    sales = [sale for sale in sales if sale.route_load_id is not None]
    payments = credit_payment_service.list_credit_payments(
        db=db,
        date_from=date_from,
        date_to=date_to,
        route_load_id=route_load_id,
    )
    payments = [
        payment for payment in payments if payment.route_load_id is not None
    ]

    workbook = create_workbook()
    add_sheet(
        workbook,
        "Cargas",
        (
            "Fecha",
            "Carga",
            "Ruta",
            "Vehículo",
            "Origen",
            "Destino",
            "Estado",
            "Producto",
            "Cantidad cargada",
            "Notas",
        ),
        (
            (
                load.load_date,
                load.load_number,
                load.route_name,
                load.vehicle_name,
                load.source_location_name,
                load.destination_location_name,
                load.status,
                item.product_name_snapshot,
                item.quantity_loaded,
                load.notes,
            )
            for load in loads
            for item in load.items
        ),
        quantity_columns=("Cantidad cargada",),
        date_columns=("Fecha",),
    )
    add_sheet(
        workbook,
        "Retornos",
        (
            "Fecha",
            "Retorno",
            "Carga",
            "Ruta",
            "Vehículo",
            "Estado",
            "Producto",
            "Cantidad retornada",
            "Notas",
        ),
        (
            (
                route_return.return_date,
                route_return.return_number,
                route_return.route_load.load_number,
                route_return.delivery_route.name,
                route_return.vehicle.name,
                route_return.status,
                item.product_name_snapshot,
                item.quantity_returned,
                route_return.notes,
            )
            for route_return in returns
            for item in route_return.items
        ),
        quantity_columns=("Cantidad retornada",),
        date_columns=("Fecha",),
    )
    add_sheet(
        workbook,
        "Cierres financieros",
        (
            "Fecha",
            "Cierre",
            "Carga",
            "Ruta",
            "Vehículo",
            "Ventas contado",
            "Ventas crédito",
            "Cobros",
            "Efectivo esperado",
            "Efectivo entregado",
            "Diferencia",
            "Estado",
            "Notas",
        ),
        (
            (
                closure.closure_date,
                closure.closure_number,
                closure.load_number_snapshot,
                closure.route_name_snapshot,
                closure.vehicle_name_snapshot,
                closure.cash_sales_total,
                closure.credit_sales_total,
                closure.payments_total,
                closure.expected_cash_total,
                closure.delivered_cash_total,
                closure.difference_amount,
                closure.status,
                closure.notes,
            )
            for closure in closures
        ),
        money_columns=(
            "Ventas contado",
            "Ventas crédito",
            "Cobros",
            "Efectivo esperado",
            "Efectivo entregado",
            "Diferencia",
        ),
        date_columns=("Fecha",),
    )
    add_sheet(
        workbook,
        "Recibos de caja",
        (
            "Fecha",
            "Recibo",
            "Cierre",
            "Carga",
            "Ruta",
            "Vehículo",
            "Efectivo entregado",
            "Efectivo recibido",
            "Diferencia",
            "Estado",
            "Notas",
        ),
        (
            (
                receipt.receipt_date,
                receipt.receipt_number,
                receipt.closure_number_snapshot,
                receipt.route_load.load_number,
                receipt.route_name_snapshot,
                receipt.vehicle_name_snapshot,
                receipt.delivered_cash_total,
                receipt.received_amount,
                receipt.difference_amount,
                receipt.status,
                receipt.notes,
            )
            for receipt in receipts
        ),
        money_columns=(
            "Efectivo entregado",
            "Efectivo recibido",
            "Diferencia",
        ),
        date_columns=("Fecha",),
    )
    activity_rows = [
        (
            sale.sale_date,
            "Venta",
            sale.route_load_number,
            sale.sale_number,
            sale.customer_name_snapshot,
            sale.total_amount,
            sale.sale_type,
            sale.status,
        )
        for sale in sales
    ]
    activity_rows.extend(
        (
            payment.payment_date,
            "Cobro",
            payment.route_load_number,
            payment.sale_number,
            payment.customer_name,
            payment.amount,
            payment.payment_method,
            payment.status,
        )
        for payment in payments
    )
    activity_rows.sort(key=lambda row: (row[0], row[3]))
    add_sheet(
        workbook,
        "Actividad comercial",
        (
            "Fecha",
            "Tipo",
            "Carga",
            "Documento",
            "Cliente",
            "Monto",
            "Detalle",
            "Estado",
        ),
        activity_rows,
        money_columns=("Monto",),
        date_columns=("Fecha",),
    )
    return workbook


def build_payroll_workbook(
    db: Session,
    *,
    date_from: date | None = None,
    date_to: date | None = None,
    employee_id: int | None = None,
    status: str | None = None,
) -> Workbook:
    runs = payroll_service.list_runs(
        db=db,
        date_from=date_from,
        date_to=date_to,
        employee_id=employee_id,
        status=status,
    )
    workbook = create_workbook()
    add_sheet(
        workbook,
        "Roles",
        (
            "Fecha rol",
            "Número",
            "Período desde",
            "Período hasta",
            "Estado",
            "Empleados",
            "Total bruto",
            "Bonos",
            "Anticipos",
            "Préstamos",
            "Otras deducciones",
            "Deducciones",
            "Neto",
            "Notas",
        ),
        (
            (
                run.payroll_date,
                run.payroll_number,
                run.period_start,
                run.period_end,
                run.status,
                run.employees_count,
                run.gross_total,
                run.bonuses_total,
                run.advances_deducted_total,
                run.loan_deductions_total,
                run.other_deductions_total,
                run.total_deductions,
                run.net_pay_total,
                run.notes,
            )
            for run in runs
        ),
        money_columns=(
            "Total bruto",
            "Bonos",
            "Anticipos",
            "Préstamos",
            "Otras deducciones",
            "Deducciones",
            "Neto",
        ),
        date_columns=("Fecha rol", "Período desde", "Período hasta"),
    )
    items = [
        item
        for run in runs
        for item in run.items
        if employee_id is None or item.employee_id == employee_id
    ]
    add_sheet(
        workbook,
        "Items de rol",
        (
            "Rol",
            "Empleado",
            "Identificación",
            "Cargo",
            "Frecuencia",
            "Sueldo base",
            "Días asistencia",
            "Ausencias",
            "Atrasos",
            "Horas",
            "Bonos",
            "Otros ingresos",
            "Bruto",
            "Anticipos",
            "Préstamos",
            "Otras deducciones",
            "Deducciones",
            "Neto",
            "Notas",
        ),
        (
            (
                item.payroll_run.payroll_number,
                item.employee_name_snapshot,
                item.employee_identification_snapshot,
                item.job_position_snapshot,
                item.payment_frequency_snapshot,
                item.base_salary_amount,
                item.attendance_days,
                item.absent_days,
                item.late_days,
                item.hours_worked,
                item.bonus_amount,
                item.other_income_amount,
                item.gross_amount,
                item.advances_deducted_amount,
                item.loan_deductions_amount,
                item.other_deductions_amount,
                item.total_deductions_amount,
                item.net_pay_amount,
                item.notes,
            )
            for item in items
        ),
        money_columns=(
            "Sueldo base",
            "Bonos",
            "Otros ingresos",
            "Bruto",
            "Anticipos",
            "Préstamos",
            "Otras deducciones",
            "Deducciones",
            "Neto",
        ),
        quantity_columns=("Horas",),
    )
    add_sheet(
        workbook,
        "Deducciones anticipos",
        (
            "Rol",
            "Empleado",
            "Anticipo",
            "Fecha anticipo",
            "Motivo",
            "Monto descontado",
        ),
        (
            (
                item.payroll_run.payroll_number,
                item.employee_name_snapshot,
                deduction.advance_number_snapshot,
                deduction.advance_date_snapshot,
                deduction.reason_snapshot,
                deduction.amount_deducted,
            )
            for item in items
            for deduction in item.advance_deductions
        ),
        money_columns=("Monto descontado",),
        date_columns=("Fecha anticipo",),
    )
    add_sheet(
        workbook,
        "Deducciones préstamos",
        (
            "Rol",
            "Empleado",
            "Préstamo",
            "Fecha préstamo",
            "Motivo",
            "Saldo de referencia",
            "Monto descontado",
        ),
        (
            (
                item.payroll_run.payroll_number,
                item.employee_name_snapshot,
                deduction.loan_number_snapshot,
                deduction.loan_date_snapshot,
                deduction.reason_snapshot,
                deduction.loan_balance_snapshot,
                deduction.amount_deducted,
            )
            for item in items
            for deduction in item.loan_deductions
        ),
        money_columns=("Saldo de referencia", "Monto descontado"),
        date_columns=("Fecha préstamo",),
    )
    return workbook


def _add_metric_sheet(
    workbook: Workbook,
    title: str,
    values: dict[str, Any],
) -> None:
    add_sheet(
        workbook,
        title,
        ("Indicador", "Valor"),
        (
            (_label(field_name), value)
            for field_name, value in values.items()
            if not isinstance(value, (list, dict))
        ),
    )


def _metric_rows(
    section: str,
    values: dict[str, Any],
) -> Iterable[tuple[str, str, Any]]:
    for field_name, value in values.items():
        if not isinstance(value, (list, dict)):
            yield section, _label(field_name), value


def _label(field_name: str) -> str:
    return LABELS.get(field_name, field_name.replace("_", " ").capitalize())


def _format_columns(
    sheet,
    headers: Sequence[str],
    selected_headers: Sequence[str],
    number_format: str,
) -> None:
    for header in selected_headers:
        if header not in headers:
            continue
        column_index = headers.index(header) + 1
        for row_index in range(2, sheet.max_row + 1):
            cell = sheet.cell(row=row_index, column=column_index)
            if cell.value is not None:
                cell.number_format = number_format


def _excel_value(value: Any) -> Any:
    if isinstance(value, Enum):
        value = value.value
    if isinstance(value, Decimal):
        return decimal_to_number(value)
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return value
    if isinstance(value, bool):
        return value
    if isinstance(value, int):
        return value
    if value is None:
        return None
    if isinstance(value, (list, tuple, set)):
        return safe_text(", ".join(str(item) for item in value))
    return safe_text(value)


def _display_length(value: Any) -> int:
    if value is None:
        return 0
    if isinstance(value, (date, datetime)):
        return 19
    return len(str(value))


def _safe_sheet_title(title: str) -> str:
    sanitized = "".join(
        "_" if character in "[]:*?/\\" else character for character in title
    ).strip()
    return (sanitized or "Reporte")[:31]


def _safe_filename(filename: str) -> str:
    sanitized = "".join(
        character
        for character in filename
        if character.isascii()
        and (character.isalnum() or character in {"-", "_", "."})
    )
    if not sanitized.endswith(".xlsx"):
        sanitized = f"{sanitized or 'reporte'}.xlsx"
    return sanitized
